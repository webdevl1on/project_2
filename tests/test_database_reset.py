"""Tiny packages and isolated SQLite targets; no production database mutation."""
import csv
from io import BytesIO
import hashlib
import json
from pathlib import Path
import sqlite3
import subprocess
import sys
from zipfile import ZipFile
import pytest
from sqlalchemy import create_engine, select, func
from sqlalchemy.orm import sessionmaker
from app.database.init_db import init_db
from app.database.maintenance import database_lease
from app.services.database_reset import reset_database
from app.services.seed_packages import HEADERS, TABLES, preflight_geography, preflight_names, check_hashes
from app.services.seed_geography import seed_geography
from app.services.seed_dictionaries import seed_name_package
from app.models import GeoUnit, GeoCityDistrict, GeoImportLink, NameDictionary, Person
from tests.test_name_cases import workbook, dictionary_data


def database(path, marker):
    with sqlite3.connect(path) as db:
        db.execute('CREATE TABLE marker (value TEXT)');db.execute('INSERT INTO marker VALUES (?)',(marker,))


def value(path):
    with sqlite3.connect(path) as db: return db.execute('SELECT value FROM marker').fetchone()[0]


def hash_manifest(directory):
    paths=[p for p in directory.rglob('*') if p.is_file() and p.name!='SHA256SUMS.txt']
    (directory/'SHA256SUMS.txt').write_text('\n'.join(hashlib.sha256(p.read_bytes()).hexdigest()+'  '+p.relative_to(directory).as_posix() for p in paths))


@pytest.fixture
def geo_package(tmp_path):
    from uuid import uuid5,NAMESPACE_URL
    directory=tmp_path/'geo';directory.mkdir();(directory/'sources').mkdir()
    (directory/'sources'/'test.txt').write_text('Synthetic source')
    source={'id':'test','version_date':'2026-01-01','local_file':'test.txt','sha256':hashlib.sha256((directory/'sources/test.txt').read_bytes()).hexdigest()}
    (directory/'sources.json').write_text(json.dumps([source]))
    ids={k:str(uuid5(NAMESPACE_URL,k)) for k in TABLES}
    contents={}
    for table in TABLES:
        row={key:'' for key in HEADERS[table].split(';')}
        row.update(id=ids[table],source_id='test',source_row='1')
        if 'country_code' in row: row['country_code']='UA'
        if table=='aliases': rows=[]
        else:
            row.update(name='Тестове',name_uk='Тестокраїна',iso_alpha2='UA',iso_alpha3='UKR',iso_numeric='804',katottg_code='UA'+str(TABLES.index(table)).zfill(17))
            for key,parent in [('country_id','countries'),('region_id','regions'),('district_id','districts'),('community_id','communities'),('settlement_id','settlements')]:
                if key in row: row[key]=ids[parent]
            rows=[row]
        contents[table]=rows
        path=directory/('streets_01.csv' if table=='streets' else table+'.csv')
        with path.open('w',encoding='utf-8-sig',newline='') as stream:
            writer=csv.DictWriter(stream,fieldnames=HEADERS[table].split(';'),delimiter=';',extrasaction='ignore');writer.writeheader();writer.writerows(rows)
    (directory/'validation.json').write_text(json.dumps({'status':'passed','counts':{k:len(v) for k,v in contents.items()}}))
    (directory/'file_validation.json').write_text(json.dumps({'status':'passed','excel_native_checks':[]}))
    hash_manifest(directory)
    return directory,ids


def test_geo_streaming_idempotency_stable_official_ids_and_city_role(geo_package,tmp_path):
    directory,ids=geo_package;prepared=preflight_geography(directory)
    engine=create_engine(f'sqlite:///{tmp_path/"geo_test.sqlite3"}');init_db(engine)
    try:
        report=seed_geography(engine,directory,prepared,batch_size=1)
        assert report['streets']['created']==1
        with engine.connect() as db:
            before=dict(db.execute(select(GeoImportLink.external_id,GeoImportLink.unit_id)).tuples().all())
        repeated=seed_geography(engine,directory,prepared,batch_size=1)
        assert sum(r['created'] for r in repeated.values())==0
        with engine.connect() as db:
            assert dict(db.execute(select(GeoImportLink.external_id,GeoImportLink.unit_id)).tuples().all())==before
            assert db.scalar(select(func.count()).select_from(GeoCityDistrict))==1
            assert db.scalar(select(func.count()).select_from(GeoUnit))==6
            assert not db.exec_driver_sql('PRAGMA foreign_key_check').fetchall()
    finally: engine.dispose()


def test_geo_failed_batch_rolls_back_whole_component(geo_package,tmp_path):
    directory,_=geo_package;engine=create_engine(f'sqlite:///{tmp_path/"rollback_test.db"}');init_db(engine)
    def fail(table,count):
        if table=='streets': raise RuntimeError('synthetic interruption')
    try:
        with pytest.raises(RuntimeError): seed_geography(engine,directory,batch_size=1,on_batch=fail)
        with engine.connect() as db: assert db.scalar(select(func.count()).select_from(GeoUnit))==0
    finally: engine.dispose()


def test_names_package_validation_and_repeat(tmp_path):
    directory=tmp_path/'names';directory.mkdir(); content=workbook([('names_male',dictionary_data())]); (directory/'project_2_dictionary_01.xlsx').write_bytes(content)
    from app.name_cases import SHEETS
    counts={k:int(k=='names_male') for k in SHEETS}
    (directory/'manifest.json').write_text(json.dumps([{'file':'project_2_dictionary_01.xlsx','records':1,'sheets':counts}]))
    (directory/'validation.json').write_text(json.dumps({'status':'passed','records':1}));hash_manifest(directory)
    prepared=preflight_names(directory);engine=create_engine(f'sqlite:///{tmp_path/"names_test.db"}');init_db(engine);factory=sessionmaker(bind=engine)
    try:
        assert seed_name_package(factory,directory,prepared)['new']==1
        assert seed_name_package(factory,directory,prepared)['unchanged']==1
        with factory() as db: assert db.scalar(select(NameDictionary)).is_verified
    finally: engine.dispose()


def test_reset_failure_preserves_existing_database_and_cleans_only_staging(tmp_path):
    target=tmp_path/'working_test.sqlite3';database(target,'original');before=target.read_bytes()
    def fail(path): database(path,'new');raise RuntimeError('failure halfway')
    with pytest.raises(RuntimeError): reset_database(target,fail)
    assert target.read_bytes()==before and value(target)=='original'
    assert not list(tmp_path.glob('*.staging-*'))


def test_publish_success_backup_and_exact_restore(tmp_path):
    target=tmp_path/'working_test.sqlite3';database(target,'original')
    report,backup=reset_database(target,lambda path:(database(path,'new') or {'ok':True}))
    assert report=={'ok':True} and value(target)=='new' and value(backup)=='original'


@pytest.mark.parametrize('suffix',['-wal','-shm','-journal'])
def test_reset_refuses_sidecars(tmp_path,suffix):
    target=tmp_path/'working_test.sqlite3';database(target,'old');Path(str(target)+suffix).touch()
    with pytest.raises(ValueError,match='WAL'): reset_database(target,lambda path:None)
    assert value(target)=='old'


def test_reset_refuses_running_application_lease(tmp_path):
    target=tmp_path/'working_test.sqlite3';database(target,'old')
    with database_lease(target,shared=True):
        with pytest.raises(ValueError,match='сервером'): reset_database(target,lambda path:None)
    assert value(target)=='old'


def test_hashes_accept_only_published_csv_newline_normalization(tmp_path):
    path=tmp_path/'test.csv';original=b'\xef\xbb\xbfkey;value\r\n"a\nb";x\r\n';path.write_bytes(original)
    hash_manifest(tmp_path);path.write_bytes(original.replace(b'\r\n',b'\n'))
    check_hashes(tmp_path)
    path.write_bytes(path.read_bytes().replace(b';x',b';changed'))
    with pytest.raises(ValueError,match='контрольна'):check_hashes(tmp_path)


def test_cli_requires_explicit_test_target_for_yes():
    from reset import main
    with pytest.raises(SystemExit):main(['--yes'])
    with pytest.raises(SystemExit):main(['--database','data/database/project2.db','--yes'])


def test_reference_workbook_positions_preserved_and_two_header_rows():
    from app.config import EXCEL_PATH
    from app.services.seed_test_people import _read_rows
    from app.services.seed_bootstrap import read_config
    from app.services.staff_import import read_positions
    assert len(read_positions(EXCEL_PATH)['positions'])==10
    assert [n for n,_ in _read_rows(EXCEL_PATH)]==[3,4,5,6]
    assert read_config(EXCEL_PATH)['valid_from'].isoformat()=='2026-01-01'


@pytest.fixture
def seed_inputs(tmp_path,geo_package):
    directory,_=geo_package
    names=tmp_path/'dictionaries';names.mkdir()
    (names/'project_2_dictionary_01.xlsx').write_bytes(workbook([('names_male',dictionary_data())]))
    from app.name_cases import SHEETS
    (names/'manifest.json').write_text(json.dumps([{'file':'project_2_dictionary_01.xlsx','records':1,'sheets':{k:int(k=='names_male') for k in SHEETS}}]))
    (names/'validation.json').write_text(json.dumps({'status':'passed','records':1}));hash_manifest(names)
    from tests.test_database_commands import sample_excel
    from openpyxl import load_workbook
    excel=sample_excel(tmp_path);book=load_workbook(excel)
    config=book.create_sheet('seed_config');config.append(['version_name','structure_code','valid_from','valid_to']);config.append(['Назва','Код','Початок','Кінець']);config.append(['Навчальний','test','2000-01-01',None])
    people=book['persons'];people.cell(1,10,'birth_country');people.cell(2,10,'Країна народження');people.cell(3,10,'Тестокраїна')
    people.cell(1,11,'registration_house');people.cell(2,11,'Будинок');people.cell(3,11,'01А/2')
    book.save(excel);book.close()
    return excel,names,directory


def test_small_full_cli_dry_run_confirmation_publish_and_backup(seed_inputs,tmp_path,capsys):
    from reset import main
    excel,names,geo=seed_inputs;target=tmp_path/'full_test.sqlite3'
    args=['--database',str(target),'--excel',str(excel),'--names',str(names),'--geography',str(geo)]
    assert main([*args,'--dry-run'])==0 and not target.exists()
    assert main(args,confirm=lambda _: 'cancel')==1 and not target.exists()
    assert main([*args,'--yes'])==0
    with sqlite3.connect(target) as db:
        assert db.execute('SELECT count(*) FROM persons').fetchone()==(2,)
        assert db.execute('SELECT count(*) FROM person_addresses').fetchone()==(1,)
        assert db.execute('SELECT count(*) FROM person_assignments').fetchone()==(0,)
    assert main([*args,'--yes'])==0
    assert len(list(tmp_path.glob('*.before-reset-*.sqlite3')))==1


def test_bad_preflight_and_failed_population_preserve_target(seed_inputs,tmp_path,monkeypatch):
    from reset import main
    from app.services import seed_bootstrap
    excel,names,geo=seed_inputs;target=tmp_path/'full_test.sqlite3';database(target,'keep')
    args=['--database',str(target),'--excel',str(excel),'--names',str(names),'--geography',str(geo),'--yes']
    def fail(*a,**kw): raise ValueError('synthetic pre-publication failure')
    monkeypatch.setattr(seed_bootstrap,'populate',fail)
    assert main(args)==2 and value(target)=='keep'
    (geo/'regions.csv').write_text('broken')
    assert main(args)==2 and value(target)=='keep'


def test_optional_assignments_have_explicit_dates_and_conflict_checks(seed_inputs):
    from openpyxl import load_workbook
    from app.services.seed_assignments import read_assignments,HEADERS
    excel,_,_=seed_inputs;book=load_workbook(excel);sheet=book.create_sheet('assignments')
    sheet.append(HEADERS);sheet.append(['РНОКПП','Посада','Початок','Кінець']);sheet.append(['0012345678',10,None,None]);book.save(excel);book.close()
    from datetime import date
    config={'valid_from':date(2000,1,1),'valid_to':None}
    with pytest.raises(ValueError): read_assignments(excel,config,{10})
    book=load_workbook(excel);book['assignments']['C3']='2000-01-01';book.save(excel);book.close()
    assert read_assignments(excel,config,{10})[0]['start']==date(2000,1,1)


def test_official_code_reuses_unit_when_external_uuid_changes(geo_package,tmp_path):
    directory,ids=geo_package;engine=create_engine(f'sqlite:///{tmp_path/"code_test.db"}');init_db(engine)
    try:
        seed_geography(engine,directory)
        with engine.connect() as db: before=db.scalar(select(GeoImportLink.unit_id).where(GeoImportLink.external_id==ids['countries']))
        from uuid import uuid4
        new_id=str(uuid4())
        for name in ('countries.csv','regions.csv'):
            p=directory/name;p.write_bytes(p.read_bytes().replace(ids['countries'].encode(),new_id.encode()))
        hash_manifest(directory);report=seed_geography(engine,directory)
        assert report['countries']['created']==0
        with engine.connect() as db:
            assert db.scalar(select(GeoImportLink.unit_id).where(GeoImportLink.external_id==new_id))==before
    finally:engine.dispose()


def test_people_and_initial_geography_roll_back_as_one_operation(tmp_path,monkeypatch):
    from tests.test_seed_test_people import isolated_db,person_row,workbook as people_workbook
    from app.services.seed_test_people import seed_test_people,REQUIRED_COLUMNS
    engine=create_engine(f'sqlite:///{tmp_path/"atomic_test.db"}');init_db(engine);factory=sessionmaker(bind=engine,autoflush=False)
    from app.models import Rank
    from app.models.reference_data.ranks import ReferenceRanks
    with factory.begin() as db:db.add(Rank(**ReferenceRanks.RANKS[0]))
    excel=people_workbook(tmp_path,[person_row()])
    import app.services.geography as geography
    def fail(*a,**kw):raise ValueError('synthetic address failure')
    from openpyxl import load_workbook
    book=load_workbook(excel);book['persons'].cell(1,len(REQUIRED_COLUMNS)+1,'registration_house');book['persons'].cell(3,len(REQUIRED_COLUMNS)+1,'01А/2');book.save(excel);book.close()
    monkeypatch.setattr(geography,'save',fail)
    report=seed_test_people(factory,excel)
    assert len(report.issues)==1 and report.created==0
    with factory() as db:assert db.scalar(select(func.count()).select_from(Person))==0
    engine.dispose()


def test_cross_book_dictionary_duplicates_are_rejected(tmp_path):
    directory=tmp_path/'names';directory.mkdir()
    from app.name_cases import SHEETS
    items=[]
    for index in (1,2):
        name=f'project_2_dictionary_0{index}.xlsx';(directory/name).write_bytes(workbook([('names_male',dictionary_data())]))
        items.append({'file':name,'records':1,'sheets':{k:int(k=='names_male') for k in SHEETS}})
    (directory/'manifest.json').write_text(json.dumps(items));(directory/'validation.json').write_text(json.dumps({'status':'passed','records':2}));hash_manifest(directory)
    with pytest.raises(ValueError,match='між книгами'):preflight_names(directory)


def test_explicit_assignment_is_saved_with_its_actual_date(seed_inputs,tmp_path):
    from openpyxl import load_workbook
    from app.services.seed_assignments import HEADERS
    from app.services.seed_bootstrap import preflight,populate
    excel,names,geo=seed_inputs;book=load_workbook(excel);sheet=book.create_sheet('assignments')
    sheet.append(HEADERS);sheet.append(['РНОКПП','Посада','Початок','Кінець']);sheet.append(['0012345678',10,'2000-01-01',None]);book.save(excel);book.close()
    prepared=preflight(excel,names,geo);target=tmp_path/'assignment_test.sqlite3'
    result=populate(target,excel,names,geo,prepared)
    assert result['person_assignments']==1
    with sqlite3.connect(target) as db:
        assert db.execute('SELECT assigned_from,assigned_to FROM person_assignments').fetchall()==[('2000-01-01',None)]


def test_unresolved_street_file_is_never_imported(geo_package,tmp_path):
    directory,_=geo_package;(directory/'streets_unresolved.csv').write_text('invalid source row deliberately not imported')
    hash_manifest(directory);prepared=preflight_geography(directory)
    engine=create_engine(f'sqlite:///{tmp_path/"unresolved_test.db"}');init_db(engine)
    try:
        report=seed_geography(engine,directory,prepared)
        assert report['streets']['rows']==1
    finally:engine.dispose()
