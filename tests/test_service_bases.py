"""Synthetic service documents and TCC data; all SQLite files are temporary."""
from datetime import date
from io import BytesIO
import sqlite3
from openpyxl import load_workbook
from sqlalchemy import select, func, event, text, create_engine
from sqlalchemy.exc import IntegrityError
import pytest
from app.database.base import Base
from app.models import Person, PersonServiceBasis, TccOffice, TccImport, PersonAssignment, PersonServiceHistory, PersonPositionMirage
from app.services import service_bases as bases, tcc_directory as directory, tcc_import as excel, dictionary_previews
from app.services.person_read import full_profile
from tests.test_name_cases import factory, client


def office(**changes):
    return {'official_name': 'Тестовий районний ТЦК та СП', 'type': 'district', 'region': 'Тестова область',
            'settlement': 'Тестове', 'is_active': '1', **changes}


def workbook(rows, change=None):
    book = load_workbook(BytesIO(excel.empty_template()))
    for row in rows:
        book.active.append([row.get(key) for key in excel.HEADERS])
    if change: change(book)
    output = BytesIO(); book.save(output); book.close()
    return output.getvalue()


def person(db):
    record = Person(); db.add(record); db.commit()
    return record.person_id


def count(db, model):
    return db.scalar(select(func.count()).select_from(model))


def test_mobilization_document_snapshot_match_and_directory_update(factory):
    with factory() as db:
        pid = person(db); oid = directory.save(db, office())
        key = bases.save(db, pid, {'service_type': 'mobilization', 'is_current': 'on', 'mobilization_date': '2022-02-24',
                                  'region': 'Тестова область', 'settlement': 'Тестове', 'document_tcc_name': office()['official_name'],
                                  'tcc_office_id': str(oid), 'document_verified': 'on'})
        record = db.get(PersonServiceBasis, key)
        assert record.match_status == 'matched' and record.document_verified
        assert record.mobilization_date == date(2022, 2, 24)
        revision = record.revision
        directory.save(db, office(official_name='Нова офіційна назва', revision=db.get(TccOffice, oid).revision), oid)
        db.refresh(record)
        assert record.document_tcc_name == office()['official_name'] and record.revision == revision
        assert full_profile(db, pid)['service_wording'] == 'за мобілізацією'
        assert count(db, PersonAssignment) == count(db, PersonServiceHistory) == count(db, PersonPositionMirage) == 0


@pytest.mark.parametrize('linked,wording,status', [(False, 'Історична назва з документа', 'unknown'), (True, 'Нестандартний військкомат', 'mismatch'), (False, '', 'unknown')])
def test_unmatched_names_and_absent_dates_are_allowed(factory, linked, wording, status):
    with factory() as db:
        pid = person(db); oid = directory.save(db, office()) if linked else None
        key = bases.save(db, pid, {'service_type': 'mobilization', 'document_tcc_name': wording, 'tcc_office_id': str(oid) if oid else ''})
        record = db.get(PersonServiceBasis, key)
        assert record.match_status == status and record.mobilization_date is None
        assert record.document_tcc_name == (wording or None)


def test_contract_has_three_independent_dates_and_history_is_explicit(factory):
    with factory() as db:
        pid = person(db)
        first = bases.save(db, pid, {'service_type': 'mobilization', 'is_current': 'on'})
        second = bases.save(db, pid, {'service_type': 'contract', 'is_current': 'on', 'contract_signed_date': '2025-01-01',
                                     'contract_start_date': '2025-02-01', 'enlistment_order_date': '2025-02-03',
                                     'enlistment_order_number': '00012/А', 'contract_unit': 'Тестова ВЧ'})
        record = db.get(PersonServiceBasis, second)
        assert (record.contract_signed_date, record.contract_start_date, record.enlistment_order_date) == (date(2025,1,1), date(2025,2,1), date(2025,2,3))
        assert not db.get(PersonServiceBasis, first).is_current
        old = db.get(PersonServiceBasis, first)
        bases.save(db, pid, {'service_type': 'mobilization', 'revision': old.revision, 'document_tcc_name': 'Виправлення документа'}, first)
        assert full_profile(db, pid)['service_wording'] == 'за контрактом'
        assert count(db, PersonServiceBasis) == 2
        bases.save(db, pid, {'service_type': 'contract', 'revision': record.revision, 'is_current': 'on', 'contract_start_date': '2025-02-01'}, second)
        assert record.contract_signed_date is None and record.enlistment_order_date is None
        assert record.contract_start_date == date(2025,2,1)


def test_current_unique_and_stale_edit_and_owner_are_protected(factory):
    with factory() as db:
        pid = person(db); other = person(db)
        key = bases.save(db, pid, {'service_type': 'contract', 'is_current': 'on'})
        record = db.get(PersonServiceBasis, key); revision = record.revision
        bases.save(db, pid, {'service_type': 'contract', 'is_current': 'on', 'revision': revision, 'contract_unit': 'Нова'}, key)
        for owner, rev in [(pid, revision), (other, record.revision)]:
            with pytest.raises(ValueError): bases.save(db, owner, {'service_type': 'mobilization', 'revision': rev}, key)
        db.add(PersonServiceBasis(person_id=pid, service_type='mobilization', is_current=True))
        with pytest.raises(IntegrityError): db.commit()
        db.rollback()
        assert count(db, PersonServiceBasis) == 1


@pytest.mark.parametrize('field', ['mobilization_date', 'contract_signed_date', 'contract_start_date', 'enlistment_order_date'])
@pytest.mark.parametrize('value', ['2026-03', '2025-02-30', '15._.2026', '20260101'])
def test_partial_or_invalid_dates_are_rejected(factory, field, value):
    with factory() as db:
        pid = person(db)
        with pytest.raises(ValueError, match='повна календарна'):
            bases.save(db, pid, {'service_type': 'mobilization' if field == 'mobilization_date' else 'contract', field: value})
        assert count(db, PersonServiceBasis) == 0


def test_date_orm_and_raw_sql_guards(factory):
    with factory() as db:
        pid = person(db); key = bases.save(db, pid, {'service_type': 'mobilization'})
        with pytest.raises(ValueError): db.get(PersonServiceBasis, key).mobilization_date = '2026-03'
        for value in ['2026-03', '2025-02-30', '0000-01-01']:
            with pytest.raises(IntegrityError):
                db.execute(text('UPDATE person_service_bases SET mobilization_date=:value WHERE id=:key'), {'key': key, 'value': value})
            db.rollback()
        assert db.get(PersonServiceBasis, key).mobilization_date is None


def test_failed_save_rolls_back_archiving_previous_current(factory):
    def fail(mapper, connection, target): raise RuntimeError('synthetic insert failure')
    with factory() as db:
        pid = person(db); key = bases.save(db, pid, {'service_type': 'mobilization', 'is_current': 'on'})
        event.listen(PersonServiceBasis, 'before_insert', fail)
        try:
            with pytest.raises(RuntimeError): bases.save(db, pid, {'service_type': 'contract', 'is_current': 'on'})
        finally: event.remove(PersonServiceBasis, 'before_insert', fail)
        assert db.get(PersonServiceBasis, key).is_current and count(db, PersonServiceBasis) == 1


def test_storage_refusal_returns_error_without_losing_current_basis(client, factory, monkeypatch):
    from sqlalchemy.exc import OperationalError
    from app.repositories.service_basis_repository import ServiceBasisRepository
    browser, _ = client
    with factory() as db:
        pid = person(db); key = bases.save(db, pid, {'service_type': 'mobilization', 'is_current': 'on'})
    def refuse(*args): raise OperationalError('synthetic', {}, Exception('refused'))
    monkeypatch.setattr(ServiceBasisRepository, 'add', refuse)
    response = browser.post(f'/persons/{pid}/service-bases', data={'service_type': 'contract', 'is_current': 'on'})
    assert response.status_code == 409 and 'Усі зміни скасовано' in response.json()['detail']
    with factory() as db:
        assert db.get(PersonServiceBasis, key).is_current and count(db, PersonServiceBasis) == 1


@pytest.mark.parametrize('field,length,kind', [('document_tcc_name', 501, 'mobilization'), ('enlistment_order_number', 101, 'contract'), ('contract_unit', 251, 'contract')])
def test_lengths(factory, field, length, kind):
    with factory() as db:
        with pytest.raises(ValueError): bases.save(db, person(db), {'service_type': kind, field: 'x'*length})


def test_directory_parent_duplicate_deactivation_and_manual_suggestions(factory):
    with factory() as db:
        first = directory.save(db, office())
        with pytest.raises(ValueError, match='уже існує'): directory.save(db, office(official_name='  ТЕСТОВИЙ РАЙОННИЙ ТЦК ТА СП '))
        child = directory.save(db, office(official_name='Тестовий відділ', type='department', parent_tcc_id=str(first)))
        with pytest.raises(ValueError, match='Циклічний'):
            directory.save(db, office(parent_tcc_id=str(child), revision=db.get(TccOffice, first).revision), first)
        with pytest.raises(ValueError, match='не існує'): directory.save(db, office(official_name='Інший', parent_tcc_id='99999'))
        directory.save(db, office(is_active=False, revision=db.get(TccOffice, first).revision), first)
        assert not directory.suggestions(db, 'районний')
        assert directory.suggestions(db, 'ВІДДІЛ', 'Тестова область')[0]['id'] == child


def test_import_repeat_controlled_update_snapshot_and_one_commit(factory):
    with factory() as db:
        content = workbook([office()]); prepared = excel.preview(db, content, 'add')
        assert prepared['counts']['new'] == 1 and not prepared['errors']
        commits = []; event.listen(db, 'after_commit', lambda db: commits.append(1))
        excel.confirm(db, prepared, 'first', reviewed=True)
        assert len(commits) == 1
        with pytest.raises(ValueError, match='уже виконано'): excel.confirm(db, prepared, 'first', reviewed=True)
        again = excel.preview(db, content, 'add'); assert again['counts']['unchanged'] == 1
        key = db.scalar(select(TccOffice.id)); pid = person(db)
        bid = bases.save(db, pid, {'service_type': 'mobilization', 'tcc_office_id': str(key), 'document_tcc_name': 'Назва за старим документом'})
        changed = workbook([office(id=key, official_name='Перейменована установа', is_active='0')])
        assert excel.preview(db, changed, 'add')['counts']['skipped'] == 1
        update = excel.preview(db, changed, 'update')
        with pytest.raises(ValueError, match='окремо'): excel.confirm(db, update, 'update', reviewed=True)
        excel.confirm(db, update, 'update', reviewed=True, confirm_updates=True)
        db.refresh(db.get(PersonServiceBasis, bid))
        assert db.get(PersonServiceBasis, bid).document_tcc_name == 'Назва за старим документом'
        assert not db.get(TccOffice, key).is_active and count(db, TccOffice) == 1


def test_import_rollback_and_stale_preview(factory):
    with factory() as db:
        prepared = excel.preview(db, workbook([office(), office(official_name='Другий ТЦК')]), 'add')
        def fail(mapper, connection, target):
            if target.official_name == 'Другий ТЦК': raise RuntimeError('second row failure')
        event.listen(TccOffice, 'before_insert', fail)
        try:
            with pytest.raises(RuntimeError): excel.confirm(db, prepared, 'failed', reviewed=True)
        finally: event.remove(TccOffice, 'before_insert', fail)
        assert count(db, TccOffice) == count(db, TccImport) == 0
        directory.save(db, office())
        with pytest.raises(ValueError, match='змінився'): excel.confirm(db, prepared, 'stale', reviewed=True)
        assert count(db, TccOffice) == 1 and count(db, TccImport) == 0


def test_excel_headers_second_row_real_row_numbers_and_duplicates(factory):
    content = workbook([office(), office(official_name='ТЕСТОВИЙ РАЙОННИЙ ТЦК ТА СП'), office(type='invented')])
    rows, errors = excel.read_workbook(content)
    assert len(rows) == 1 and 'рядок 4' in errors[0] and 'рядок 5' in errors[1]
    assert excel.read_workbook(excel.empty_template())[1] == ['Немає заповнених рядків.']
    with pytest.raises(ValueError, match='Рядок 1'): excel.read_workbook(workbook([], lambda b: setattr(b.active['A1'], 'value', 'wrong')))
    with pytest.raises(ValueError, match='аркуша'): excel.read_workbook(workbook([], lambda b: setattr(b.active, 'title', 'wrong')))
    with pytest.raises(ValueError): excel.read_workbook(b'not excel')
    with pytest.raises(ValueError): excel.read_workbook(b'x'*(5*1024*1024+1))
    assert 'Формули' in excel.read_workbook(workbook([office(official_name='=1+1')]))[1][0]
    assert 'id' in excel.read_workbook(workbook([office(id=1.5)]))[1][0]


def test_preview_owner_resource_and_parent_cycles(factory):
    with factory() as db:
        first = directory.save(db, office()); second = directory.save(db, office(official_name='Другий'))
        prepared = excel.preview(db, workbook([office(id=first, parent_tcc_id=second), office(id=second, official_name='Другий', parent_tcc_id=first)]), 'update')
        assert any('Циклічний' in error for error in prepared['errors'])
        prepared = excel.preview(db, workbook([office(official_name='Третій', parent_tcc_id=99999)]), 'add')
        assert any('не існує' in error for error in prepared['errors'])
        prepared = excel.preview(db, workbook([office(official_name='Третій')]), 'add')
        token = dictionary_previews.store(prepared, 'admin-one', resource='tcc')
        with pytest.raises(ValueError): dictionary_previews.load(token, 'admin-two', resource='tcc')
        with pytest.raises(ValueError): dictionary_previews.load(token, 'admin-one')
        assert dictionary_previews.load(token, 'admin-one', resource='tcc')[0] == prepared


def test_http_admin_protection_csrf_preview_confirm_and_profile(client, factory):
    browser, secret = client
    for path in ['/admin/tcc', '/admin/tcc/template.xlsx', '/admin/tcc/new']:
        assert browser.get(path, follow_redirects=False).status_code == 401
    assert browser.post('/admin/tcc/import/preview', files={'file': ('tcc.xlsx', workbook([office()]))}).status_code == 401
    assert browser.post('/admin/tcc/new', data=office()).status_code == 401
    browser.post('/admin/login', data={'token': secret})
    from lxml import html
    csrf = html.fromstring(browser.get('/admin').text).xpath('//input[@name="csrf"]/@value')[0]
    assert browser.post('/admin/tcc/new', data=office()).status_code == 403
    response = browser.post('/admin/tcc/import/preview', data={'csrf': csrf}, files={'file': ('tcc.xlsx', workbook([office()]))})
    assert response.status_code == 200
    token = html.fromstring(response.text).xpath('//input[@name="token"]/@value')[0]
    assert browser.post('/admin/tcc/import/confirm', data={'csrf': csrf, 'token': token, 'reviewed': 'on'}).status_code == 200
    assert browser.post('/admin/tcc/import/confirm', data={'csrf': csrf, 'token': token, 'reviewed': 'on'}).status_code == 409
    with factory() as db:
        pid = person(db)
        db.add(PersonServiceHistory(person_id=pid, service_type='contract', is_active=True)); db.commit()
    assert 'Вид служби: Не зазначено' in browser.get(f'/person_full_info/{pid}').text
    response = browser.post(f'/persons/{pid}/service-bases', data={'service_type': 'contract', 'is_current': 'on', 'contract_start_date': '2025-03'})
    assert response.status_code == 409 and 'повна календарна' in response.json()['detail']
    assert browser.post(f'/persons/{pid}/service-bases', data={'service_type': 'contract', 'is_current': 'on'}).status_code == 200
    assert 'Вид служби: за контрактом' in browser.get(f'/person_full_info/{pid}').text


def test_admin_edit_error_preserves_values_and_import_update_rechecks_revision(client, factory):
    from tests.test_name_cases import login
    browser, secret = client; csrf = login(browser, secret)
    with factory() as db:
        key = directory.save(db, office()); revision = db.get(TccOffice, key).revision
        prepared = excel.preview(db, workbook([office(id=key, official_name='Назва із preview')]), 'update')
        directory.save(db, office(revision=revision, official_name='Змінено паралельно'), key)
        with pytest.raises(ValueError, match='змінився'):
            excel.confirm(db, prepared, 'stale-update', reviewed=True, confirm_updates=True)
    response = browser.post(f'/admin/tcc/{key}/edit', data={**office(official_name='Моя незбережена назва'), 'revision': revision, 'csrf': csrf})
    assert response.status_code == 409 and 'Моя незбережена назва' in response.text
    assert 'змінено після відкриття' in response.text


@pytest.mark.parametrize('value', ['-1', 'invalid', '9'*50])
def test_invalid_reference_ids_have_readable_errors(factory, value):
    with factory() as db:
        pid = person(db)
        with pytest.raises(ValueError, match='ID'):
            bases.save(db, pid, {'service_type': 'mobilization', 'tcc_office_id': value})
        with pytest.raises(ValueError, match='ID'): directory.save(db, office(parent_tcc_id=value))


@pytest.mark.parametrize('key', [0, -1, 10**30])
def test_invalid_edit_ids_cannot_create_new_records(factory, key):
    with factory() as db:
        pid = person(db)
        with pytest.raises(ValueError): bases.save(db, pid, {'service_type': 'contract'}, key)
        with pytest.raises(ValueError): directory.save(db, office(), key)
        assert count(db, PersonServiceBasis) == count(db, TccOffice) == 0


def old_database(path):
    from app.models.service_basis import NEW_TABLE_NAMES
    engine = create_engine(f'sqlite:///{path}')
    Base.metadata.create_all(engine, tables=[table for table in Base.metadata.sorted_tables if table.name not in NEW_TABLE_NAMES])
    from sqlalchemy.orm import Session
    with Session(engine) as db:
        db.add(Person(tax_number='0000000001')); db.commit()
    engine.dispose()


def test_migration_backup_preserves_all_data_triggers_and_is_idempotent(tmp_path):
    from app.database.migrate_service_bases import migrate
    from app.models.service_basis import NEW_TABLE_NAMES
    path = tmp_path/'before.sqlite3'; old_database(path)
    with sqlite3.connect(path) as connection:
        prior = connection.execute('SELECT * FROM persons').fetchall()
        triggers = connection.execute("SELECT name,sql FROM sqlite_master WHERE type='trigger'").fetchall()
    backup = migrate(path)
    assert backup.is_file() and backup.stat().st_mode & 0o777 == 0o600
    with sqlite3.connect(path) as connection:
        assert connection.execute('SELECT * FROM persons').fetchall() == prior
        assert set(triggers) <= set(connection.execute("SELECT name,sql FROM sqlite_master WHERE type='trigger'").fetchall())
        assert connection.execute('PRAGMA foreign_key_check').fetchall() == []
        for table in NEW_TABLE_NAMES: assert connection.execute(f'SELECT count(*) FROM {table}').fetchone() == (0,)
    with sqlite3.connect(backup) as connection: assert connection.execute('SELECT * FROM persons').fetchall() == prior
    assert migrate(path) is None


def test_migration_ddl_failure_rolls_back_and_keeps_backup(tmp_path, monkeypatch):
    from app.database import migrate_service_bases as migration
    from app.models.service_basis import NEW_TABLE_NAMES
    path = tmp_path/'rollback.sqlite3'; old_database(path)
    original = migration.create_table
    def fail(connection, table):
        original(connection, table)
        if table.name == 'tcc_offices': raise RuntimeError('synthetic DDL refusal')
    monkeypatch.setattr(migration, 'create_table', fail)
    with pytest.raises(RuntimeError): migration.migrate(path)
    with sqlite3.connect(path) as connection:
        present = {r[0] for r in connection.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        assert not NEW_TABLE_NAMES & present
        assert connection.execute('SELECT tax_number FROM persons').fetchone() == ('0000000001',)
    assert len(list(tmp_path.glob('*.before-service-bases-*.sqlite3'))) == 1


def test_migration_rejects_incomplete_schema_and_running_server_lease(tmp_path):
    from app.database.migrate_service_bases import migrate
    from app.database.maintenance import database_lease
    from app.database.schema_check import require_current_schema
    path = tmp_path/'lease.sqlite3'; old_database(path)
    engine = create_engine(f'sqlite:///{path}')
    with pytest.raises(ValueError, match='migrate_service_bases'): require_current_schema(engine)
    engine.dispose()
    with database_lease(path, shared=True):
        with pytest.raises(ValueError, match='використовується'): migrate(path)
    with sqlite3.connect(path) as connection:
        connection.execute('DROP TRIGGER guard_person_assignments_insert')
    with pytest.raises(ValueError, match='тригерів'): migrate(path)
    assert not list(tmp_path.glob('*.before-service-bases-*.sqlite3'))


def test_generated_and_committed_excel_templates_have_same_empty_structure():
    from pathlib import Path
    for content in [excel.empty_template(), (Path(__file__).resolve().parents[1]/'docs/templates/tcc_offices_empty.xlsx').read_bytes()]:
        book = load_workbook(BytesIO(content))
        try:
            assert book.sheetnames == ['tcc_offices'] and book.active.max_row == 2
            assert tuple(cell.value for cell in book.active[1]) == excel.HEADERS
            assert tuple(cell.value for cell in book.active[2]) == excel.LABELS
        finally: book.close()
