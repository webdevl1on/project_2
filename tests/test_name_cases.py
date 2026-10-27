"""Synthetic words and isolated SQLite only; no production dictionary or people."""
from io import BytesIO
import secrets
from lxml import html
from openpyxl import load_workbook
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select, func, event
from sqlalchemy.orm import sessionmaker
from app.database.base import Base
from app.database.db import get_db
from app.models import (Person, PersonName, PersonRank, PersonNameCases, NameDictionary,
                        NameDictionaryImport, Rank, EventType, PersonEvent)
from app.models.reference_data.ranks import ReferenceRanks
from app.name_cases import CASES, SHEETS, EXCEL_HEADERS, EXCEL_LABELS
from app.services import name_cases, name_dictionary_import as excel, dictionary_previews
from app.services.person_registration import register_person
from app.services.service_person_create import ServicePersonCreate
from app.services.service_person_name import ServicePersonNameFunc
from app.validators.name_cases import search_key


@pytest.fixture
def factory(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'names.sqlite3'}")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False)
    with factory.begin() as db:
        db.add(Rank(**ReferenceRanks.RANKS[0]))
    yield factory
    engine.dispose()


def person_data(**changes):
    return {"sex": "male", "surname": "Тестенко", "name": "Тестослав", "patronymic": "Тестович",
            "tax_number": "0123456789", "birth_month": "3", "birth_year": "1993", "rank_id": "1", **changes}


def dictionary_data(kind="name", sex="male", base="Тестослав", **changes):
    # Repeated forms are deliberate fixture values, not a declension algorithm.
    return {"kind": kind, "sex": sex, "base_form": base, **{case: base for case in CASES}, **changes}


def workbook(rows=(), *, change=None):
    book = load_workbook(BytesIO(excel.empty_template()))
    for sheet, row in rows:
        book[sheet].append([row.get(key) for key in EXCEL_HEADERS])
    if change:
        change(book)
    target = BytesIO(); book.save(target); book.close()
    return target.getvalue()


def count(db, model):
    return db.scalar(select(func.count()).select_from(model))


def test_verified_lookup_is_case_preserving_gender_specific_and_does_not_guess(factory):
    with factory() as db:
        male = dictionary_data(base="Сокіл", genitive="Сокола")
        female = dictionary_data(sex="female", base="Сокіл", genitive="Сокіл")
        name_cases.save_dictionary(db, male, verified=True)
        name_cases.save_dictionary(db, female, verified=True)
        assert name_cases.lookup(db, "name", "СОКІЛ", "male")["forms"]["genitive"] == "Сокола"
        assert name_cases.lookup(db, "name", "сокіл", "female")["forms"]["genitive"] == "Сокіл"
        assert not name_cases.lookup(db, "name", "Невідоме", "male")["found"]
        with pytest.raises(ValueError):
            name_cases.lookup(db, "name", "Сокіл", "")
        draft = name_cases.save_dictionary(db, dictionary_data(base="О’Ніл"), verified=False)
        assert not name_cases.lookup(db, "name", "О'НІЛ", "male")["found"]
        assert db.get(NameDictionary, draft).base_form == "О’Ніл"
        assert search_key("  О’Ніл ") == search_key("о'ніл")
        with pytest.raises(ValueError, match="вже існує"):
            name_cases.save_dictionary(db, dictionary_data(base="о'ніл"), verified=True)


@pytest.mark.parametrize("gender", [None, "", "unknown", "MALE"])
def test_new_people_require_explicit_sex_but_legacy_nullable_rows_exist(factory, gender):
    with factory() as db:
        db.add(Person()); db.commit()
        assert db.scalar(select(Person)).sex is None
        with pytest.raises(ValueError, match="стать"):
            ServicePersonCreate.create_person(person_data(sex=gender), db)
        assert count(db, Person) == 1


def test_registration_saves_partial_dates_cases_and_initial_rank_atomically(factory):
    with factory() as db:
        commits = []
        event.listen(db, "after_commit", lambda db: commits.append(True))
        person = register_person(db, person_data(name_genitive="Тестослава", name_verified="on"))
        assert len(commits) == 1
        assert person.sex == "male" and person.tax_number == "0123456789"
        assert (person.birth_day, person.birth_month, person.birth_year) == (None, 3, 1993)
        assert person.names[0].name == "Тестослав"
        saved = db.scalar(select(PersonNameCases).where(PersonNameCases.kind == "name"))
        assert saved.person_name_id == person.names[0].person_name_id
        assert saved.genitive == "Тестослава" and saved.dative is None and not saved.is_verified
        rank = person.ranks[0]
        assert rank.is_active and rank.rank_order_number is None and rank.rank_order_date is None
        assert (rank.rank_order_day, rank.rank_order_month, rank.rank_order_year) == (None, None, None)
        assert count(db, NameDictionary) == 0


@pytest.mark.parametrize("changes", [{"rank_id": "99999"}, {"name_bound_base": "Інше", "name_genitive": "Тестослава"},
                                     {"name_bound_sex": "female", "name_genitive": "Тестослава"}])
def test_registration_failure_does_not_leave_any_partial_records(factory, changes):
    with factory() as db:
        with pytest.raises(ValueError):
            register_person(db, person_data(**changes))
        assert [count(db, model) for model in (Person, PersonName, PersonNameCases, PersonRank)] == [0, 0, 0, 0]


def test_database_refusal_rolls_back_person_name_cases_and_rank(factory):
    def refuse(mapper, connection, target):
        raise RuntimeError("synthetic rank write refusal")
    event.listen(PersonRank, "before_insert", refuse)
    try:
        with factory() as db:
            with pytest.raises(RuntimeError):
                register_person(db, person_data(name_genitive="Тестослава"))
            assert [count(db, model) for model in (Person, PersonName, PersonNameCases, PersonRank)] == [0, 0, 0, 0]
    finally:
        event.remove(PersonRank, "before_insert", refuse)


def test_personal_priority_history_and_stale_snapshot_protection(factory):
    with factory() as db:
        key = name_cases.save_dictionary(db, dictionary_data(genitive="Словникового", dative="Словниковому"), verified=True)
        person = register_person(db, person_data(name_genitive="Особистого"))
        initial = person.names[0]
        resolved = name_cases.resolve(db, initial, person.sex)["name"]
        assert resolved["forms"]["genitive"] == "Особистого" and resolved["sources"]["genitive"] == "personal"
        assert resolved["forms"]["dative"] == "Словниковому" and resolved["sources"]["dative"] == "dictionary"
        entry = db.get(NameDictionary, key)
        name_cases.save_dictionary(db, dictionary_data(genitive="Оновленого"), key=key, revision=entry.revision, verified=True)
        assert name_cases.resolve(db, initial, person.sex)["name"]["forms"]["genitive"] == "Особистого"
        new_name = ServicePersonNameFunc.add_new_name(person.person_id, {"surname": "Тестенко", "name": "Тестослав"}, db)
        assert name_cases.resolve(db, new_name, person.sex)["name"]["forms"]["genitive"] == "Оновленого"
        assert name_cases.resolve(db, initial, "female")["name"]["stale"]
        ServicePersonNameFunc.update_person_name(initial.person_name_id, {"name": "Інший"}, db)
        after = name_cases.resolve(db, initial, person.sex)["name"]
        assert after["stale"] and after["forms"]["genitive"] is None
        assert any("повторна" in text for text in after["warnings"])


def test_template_is_empty_and_documents_exact_headers():
    book = load_workbook(BytesIO(excel.empty_template()))
    assert book.sheetnames == list(SHEETS)
    for sheet in book:
        assert sheet.max_row == 2
        assert tuple(cell.value for cell in sheet[1]) == EXCEL_HEADERS
        assert tuple(cell.value for cell in sheet[2]) == EXCEL_LABELS
        assert sheet.freeze_panes == "A3"
    book.close()


def test_all_six_sheets_derive_kind_and_sex_and_skip_labels(factory):
    content = workbook([(sheet, dictionary_data(sex="ignored", kind="ignored")) for sheet in SHEETS])
    with factory() as db:
        prepared = excel.preview(db, content, "add")
        assert prepared["errors"] == [] and prepared["counts"]["new"] == 6
        assert [(row["sheet"], row["row"], row["values"]["kind"], row["values"]["sex"])
                for row in prepared["rows"]] == [(sheet, 3, kind, sex) for sheet, (kind, sex) in SHEETS.items()]
        excel.confirm(db, prepared, "six-sheets", reviewed=True)
        assert count(db, NameDictionary) == 6
        assert excel.preview(db, content, "add")["counts"]["unchanged"] == 6


def test_second_row_is_not_data_and_errors_use_excel_row_numbers():
    def change(book):
        sheet = book["names_male"]
        sheet["A2"] = "=not_dictionary_data"
        sheet.append([None] * len(EXCEL_HEADERS))
        sheet.append(["Помилка", "Інше", *(["Помилка"] * 6)])
    rows, errors = excel.read_workbook(workbook([("names_male", dictionary_data())], change=change))
    assert len(rows) == 1 and rows[0]["row"] == 3
    assert len(errors) == 1 and "names_male, рядок 5" in errors[0]


def test_excel_rejects_old_sex_column_and_template_artifact_matches_contract():
    from pathlib import Path
    rows, errors = excel.read_workbook(workbook(change=lambda book: book["names_male"].insert_cols(2)))
    assert not rows and any("рядок 1" in error for error in errors)
    artifact = Path(__file__).resolve().parents[1] / "docs/templates/name_dictionaries_empty.xlsx"
    book = load_workbook(artifact)
    assert book.sheetnames == list(SHEETS)
    for sheet in book:
        assert sheet.max_row == 2
        assert tuple(cell.value for cell in sheet[1]) == EXCEL_HEADERS
        assert tuple(cell.value for cell in sheet[2]) == EXCEL_LABELS
    book.close()


@pytest.mark.parametrize("change,message", [
    (lambda book: book.remove(book["names_male"]), "Відсутні аркуші"),
    (lambda book: setattr(book["names_male"]["A1"], "value", "wrong"), "заголовки"),
    (lambda book: setattr(book["names_male"]["I3"], "value", "unexpected"), "межами"),
    (lambda book: setattr(book["names_male"]["C3"], "value", None), "Родовий"),
    (lambda book: setattr(book["names_male"]["C3"], "value", '=A3'), "Формули"),
    (lambda book: setattr(book["names_male"]["B3"], "value", "Інше"), "повинні збігатися"),
])
def test_import_errors_never_write(factory, change, message):
    with factory() as db:
        content = workbook([("names_male", dictionary_data())], change=change)
        try:
            prepared = excel.preview(db, content, "add")
        except ValueError as error:
            assert message in str(error)
        else:
            assert any(message in error for error in prepared["errors"])
            with pytest.raises(ValueError):
                excel.confirm(db, prepared, "bad", reviewed=True)
        assert count(db, NameDictionary) == 0


def test_duplicate_excel_keys_are_errors_even_with_different_case(factory):
    with factory() as db:
        content = workbook([("names_male", dictionary_data()), ("names_male", dictionary_data(base="ТЕСТОСЛАВ"))])
        prepared = excel.preview(db, content, "add")
        assert prepared["counts"]["errors"] == 1
        assert "names_male, рядок 4" in prepared["errors"][0]
        with pytest.raises(ValueError):
            excel.confirm(db, prepared, "duplicate", reviewed=True)
        assert count(db, NameDictionary) == 0


def test_import_modes_confirmation_replay_and_personal_forms_are_independent(factory):
    with factory() as db:
        initial = workbook([("names_male", dictionary_data(genitive="Першого"))])
        prepared = excel.preview(db, initial, "add")
        assert prepared["counts"]["new"] == 1 and count(db, NameDictionary) == 0
        commits = []; event.listen(db, "after_commit", lambda db: commits.append(True))
        excel.confirm(db, prepared, "first", reviewed=True)
        assert len(commits) == 1
        with pytest.raises(ValueError, match="уже підтверджено"):
            excel.confirm(db, prepared, "first", reviewed=True)
        person = register_person(db, person_data(name_genitive="Персонального"))
        personal = db.scalar(select(PersonNameCases).where(PersonNameCases.kind == "name"))
        snapshot_revision = personal.revision
        changed = workbook([("names_male", dictionary_data(genitive="Другого"))])
        only_new = excel.preview(db, changed, "add")
        assert only_new["counts"]["updated"] == 0 and only_new["counts"]["skipped"] == 1
        update = excel.preview(db, changed, "update")
        assert update["counts"]["updated"] == 1
        with pytest.raises(ValueError, match="Окремо"):
            excel.confirm(db, update, "second", reviewed=True)
        excel.confirm(db, update, "second", reviewed=True, confirm_updates=True)
        db.refresh(personal)
        assert personal.genitive == "Персонального" and personal.revision == snapshot_revision
        unchanged = excel.preview(db, changed, "update")
        assert unchanged["counts"]["unchanged"] == 1 and unchanged["counts"]["updated"] == 0
        assert name_cases.resolve(db, person.names[0], "male")["name"]["forms"]["genitive"] == "Персонального"


def test_import_failure_rolls_back_all_rows_and_receipt(factory):
    def refuse(mapper, connection, target):
        if target.kind == "surname":
            raise RuntimeError("synthetic second dictionary failure")
    event.listen(NameDictionary, "before_insert", refuse)
    try:
        with factory() as db:
            content = workbook([("names_male", dictionary_data()), ("surname_male", dictionary_data(kind="surname", base="Тестенко"))])
            # Sheet iteration order puts surnames first; force failure on the second write.
            def fail_second(mapper, connection, target):
                if target.kind == "name": raise RuntimeError("second write")
            event.remove(NameDictionary, "before_insert", refuse)
            event.listen(NameDictionary, "before_insert", fail_second)
            try:
                with pytest.raises(RuntimeError):
                    excel.confirm(db, excel.preview(db, content, "add"), "refused", reviewed=True)
            finally:
                event.remove(NameDictionary, "before_insert", fail_second)
            assert count(db, NameDictionary) == count(db, NameDictionaryImport) == 0
    finally:
        if event.contains(NameDictionary, "before_insert", refuse):
            event.remove(NameDictionary, "before_insert", refuse)


def test_preview_rejects_changes_since_review_and_owner_mismatch(factory):
    with factory() as db:
        prepared = excel.preview(db, workbook([("names_male", dictionary_data())]), "add")
        token = dictionary_previews.store(prepared, "first-session")
        with pytest.raises(ValueError, match="іншій"):
            dictionary_previews.load(token, "second-session")
        name_cases.save_dictionary(db, dictionary_data(), verified=True)
        with pytest.raises(ValueError, match="змінився"):
            excel.confirm(db, prepared, "stale", reviewed=True)
        assert count(db, NameDictionary) == 1 and count(db, NameDictionaryImport) == 0


@pytest.fixture
def client(factory, monkeypatch):
    from app.main import app
    token = secrets.token_urlsafe(32)
    monkeypatch.setenv("PROJECT2_ADMIN_TOKEN", token)
    monkeypatch.delenv("PROJECT2_ADMIN_ALLOW_HTTP", raising=False)
    monkeypatch.delenv("PROJECT2_DEV_ADMIN_OPEN", raising=False)
    def override():
        with factory() as db: yield db
    old = app.dependency_overrides.copy(); app.dependency_overrides[get_db] = override
    try:
        with TestClient(app, base_url="https://testserver") as client:
            yield client, token
    finally:
        app.dependency_overrides.clear(); app.dependency_overrides.update(old)


def login(client, token):
    assert client.post("/admin/login", data={"token": token}, follow_redirects=False).status_code == 303
    page = client.get("/admin")
    return html.fromstring(page.text).xpath('//input[@name="csrf"]/@value')[0]


def test_admin_denies_mutations_without_login_or_csrf_and_is_disabled_by_default(client, monkeypatch):
    client, token = client
    for path in ("/admin/dictionaries/new", "/admin/dictionaries/import/preview", "/admin/dictionaries/import/confirm", "/admin/dictionaries/1/edit"):
        assert client.post(path, data={}).status_code == 401
    assert client.get("/admin/dictionaries/template.xlsx").status_code == 401
    assert client.post("/admin/login", data={"token": "wrong"}).status_code == 401
    csrf = login(client, token)
    assert client.post("/admin/dictionaries/new", data=dictionary_data()).status_code == 403
    assert client.post("/admin/dictionaries/new", data={**dictionary_data(), "csrf": "чужий"}).status_code == 403
    assert client.post("/admin/logout", data={"csrf": csrf}, follow_redirects=False).status_code == 303
    monkeypatch.delenv("PROJECT2_ADMIN_TOKEN")
    assert client.post("/admin/dictionaries/new", data={}).status_code == 503


def test_admin_cookie_transport_and_manual_editor(client):
    client, token = client
    response = client.post("/admin/login", data={"token": token}, follow_redirects=False)
    cookie = response.headers["set-cookie"]
    assert "HttpOnly" in cookie and "Secure" in cookie and "SameSite=strict" in cookie and "Path=/admin" in cookie
    csrf = html.fromstring(client.get("/admin").text).xpath('//input[@name="csrf"]/@value')[0]
    response = client.post("/admin/dictionaries/new", data={**dictionary_data(), "csrf": csrf, "is_verified": "on"})
    assert response.status_code == 200 and "Тестослав" in response.text
    edit_url = html.fromstring(response.text).xpath('//a[contains(@href,"/edit")]/@href')[0]
    page = client.get(edit_url)
    revision = html.fromstring(page.text).xpath('//input[@name="revision"]/@value')[0]
    edited = client.post(edit_url, data={**dictionary_data(genitive="Тестослава"), "csrf": csrf, "revision": revision, "is_verified": "on"})
    assert edited.status_code == 200
    assert client.get("/name-dictionaries/lookup", params={"kind":"name","sex":"male","word":"Тестослав"}).json()["forms"]["genitive"] == "Тестослава"
    with TestClient(client.app, base_url="http://testserver") as plain:
        assert plain.get("/admin/login").status_code == 403
        assert plain.post("/admin/login", data={"token": token}).status_code == 403


def test_admin_excel_preview_confirm_and_replay(client, factory):
    client, token = client; csrf = login(client, token)
    assert client.get("/admin/dictionaries/template.xlsx").content[:2] == b"PK"
    content = workbook([("names_male", dictionary_data())])
    preview = client.post("/admin/dictionaries/import/preview", files={"file": ("names.xlsx", content)}, data={"mode": "add", "csrf": csrf})
    assert preview.status_code == 200 and 'data-count="new">1' in preview.text
    preview_token = html.fromstring(preview.text).xpath('//input[@name="token"]/@value')[0]
    with factory() as db: assert count(db, NameDictionary) == 0
    refused = client.post("/admin/dictionaries/import/confirm", data={"csrf": csrf, "token": preview_token})
    assert refused.status_code == 409
    saved = client.post("/admin/dictionaries/import/confirm", data={"csrf": csrf, "token": preview_token, "reviewed": "true"})
    assert saved.status_code == 200 and "Імпорт завершено" in saved.text
    assert client.post("/admin/dictionaries/import/confirm", data={"csrf": csrf, "token": preview_token, "reviewed": "true"}).status_code == 409
    with factory() as db: assert count(db, NameDictionary) == 1


def test_creation_error_retains_all_fields_and_success_has_two_destinations(client, factory):
    client, _ = client
    values = person_data(birth_day="31", birth_month="2", surname_genitive="Тестенка", name_dative="Тестославу")
    response = client.post("/persons/new", data=values)
    assert response.status_code == 400
    tree = html.fromstring(response.text)
    for key in ("surname_genitive", "name_dative", "birth_day", "birth_month", "tax_number"):
        assert tree.xpath(f'//input[@name="{key}"]/@value')[0] == values[key]
    assert tree.xpath('//input[@name="sex" and @checked]/@value') == ["male"]
    assert tree.xpath('//option[@selected]/@value') == ["1"]
    success = client.post("/persons/new", data=person_data(name_genitive="Тестослава"))
    assert success.status_code == 200 and "Людину збережено" in success.text
    tree = html.fromstring(success.text)
    assert tree.xpath('//a[@href="/persons"]') and tree.xpath('//a[starts-with(@href,"/person_full_info/")]')
    with factory() as db:
        assert count(db, Person) == 1 and count(db, PersonRank) == 1


def test_person_names_keep_case_apostrophes_hyphens_and_short_names(factory):
    with factory() as db:
        person = register_person(db, person_data(surname="О’Ніл", name="Ян", patronymic="Анна-Марія"))
        assert (person.names[0].surname, person.names[0].name, person.names[0].patronymic) == ("О’Ніл", "Ян", "Анна-Марія")


def test_admin_editor_stale_revision_retains_draft_and_does_not_overwrite(client, factory):
    client, token = client; csrf = login(client, token)
    with factory() as db:
        key = name_cases.save_dictionary(db, dictionary_data(), verified=True)
    page = client.get(f"/admin/dictionaries/{key}/edit")
    revision = html.fromstring(page.text).xpath('//input[@name="revision"]/@value')[0]
    with factory() as db:
        name_cases.save_dictionary(db, dictionary_data(genitive="Актуального"), key=key, revision=int(revision), verified=True)
    response = client.post(f"/admin/dictionaries/{key}/edit", data={**dictionary_data(genitive="Чернеткового"),
        "revision":revision,"csrf":csrf,"is_verified":"on"})
    assert response.status_code == 409
    assert html.fromstring(response.text).xpath('//input[@name="genitive"]/@value') == ["Чернеткового"]
    with factory() as db:
        assert db.get(NameDictionary, key).genitive == "Актуального"


def test_admin_session_tampering_expiry_and_token_rotation_are_rejected(client, monkeypatch):
    client, token = client
    login(client, token)
    client.cookies.set("project2_admin", "tampered", domain="testserver.local", path="/admin")
    assert client.get("/admin/dictionaries").status_code == 401
    client.cookies.clear()
    import app.security.admin as admin
    expires = str(int(admin.time.time()) - 1)
    client.cookies.set("project2_admin", admin.signed(expires+":"+secrets.token_hex(32), "session"),
                       domain="testserver.local", path="/admin")
    assert client.get("/admin/dictionaries").status_code == 401
    client.cookies.clear(); login(client, token)
    monkeypatch.setenv("PROJECT2_ADMIN_TOKEN", secrets.token_urlsafe(32))
    assert client.get("/admin/dictionaries").status_code == 401


def test_dev_admin_local_all_operations_without_credentials(client, factory, monkeypatch):
    protected, _ = client
    monkeypatch.setenv("PROJECT2_DEV_ADMIN_OPEN", "1")
    monkeypatch.delenv("PROJECT2_ADMIN_TOKEN")
    with TestClient(protected.app, base_url="http://localhost", client=("127.0.0.1", 50000)) as local:
        assert local.get("/admin").status_code == 200
        assert local.get("/admin/login", follow_redirects=False).headers["location"] == "/admin"
        assert local.post("/admin/login", data={}).status_code == 200
        assert not local.cookies
        assert local.get("/admin/dictionaries/template.xlsx").content[:2] == b"PK"
        assert local.get("/admin/dictionaries/new").status_code == 200
        created = local.post("/admin/dictionaries/new", data=dictionary_data())
        assert created.status_code == 200
        edit_url = html.fromstring(created.text).xpath('//a[contains(@href,"/edit")]/@href')[0]
        revision = html.fromstring(local.get(edit_url).text).xpath('//input[@name="revision"]/@value')[0]
        assert local.post(edit_url, data={**dictionary_data(), "revision": revision}).status_code == 200
        content = workbook([("names_male", dictionary_data(base="Іншослав"))])
        preview = local.post("/admin/dictionaries/import/preview", files={"file": ("names.xlsx", content)})
        assert preview.status_code == 200
        preview_token = html.fromstring(preview.text).xpath('//input[@name="token"]/@value')[0]
        saved = local.post("/admin/dictionaries/import/confirm", data={"token": preview_token, "reviewed": "true"})
        assert saved.status_code == 200 and "Імпорт завершено" in saved.text
        with factory() as db:
            assert count(db, NameDictionary) == 2
        assert local.post("/admin/logout").status_code == 200
        monkeypatch.delenv("PROJECT2_DEV_ADMIN_OPEN")
        assert local.get("/admin").status_code == 503
        assert local.post("/admin/dictionaries/new", data={}).status_code == 503


@pytest.mark.parametrize("host,peer,headers", [
    ("localhost", "192.168.1.10", {}),
    ("localhost", "203.0.113.10", {}),
    ("192.168.1.5", "127.0.0.1", {}),
    ("203.0.113.5", "127.0.0.1", {}),
    ("example.com", "127.0.0.1", {}),
    ("localhost", "127.0.0.1", {"X-Forwarded-For": "127.0.0.1"}),
    ("localhost", "127.0.0.1", {"Forwarded": "for=127.0.0.1"}),
    ("localhost", "127.0.0.1", {"X-Real-IP": "127.0.0.1"}),
    ("localhost", "127.0.0.1", {"Origin": "https://example.com"}),
    ("localhost", "127.0.0.1", {"Sec-Fetch-Site": "cross-site"}),
])
def test_dev_admin_rejects_nonlocal_and_forwarded_requests(client, monkeypatch, host, peer, headers):
    protected, _ = client
    monkeypatch.setenv("PROJECT2_DEV_ADMIN_OPEN", "1")
    monkeypatch.delenv("PROJECT2_ADMIN_TOKEN")
    with TestClient(protected.app, base_url="http://" + host, client=(peer, 50000), headers=headers) as remote:
        assert remote.get("/admin").status_code == 503
        assert remote.get("/admin/dictionaries/template.xlsx").status_code == 503
        assert remote.post("/admin/dictionaries/new", data={}).status_code == 503
        assert remote.post("/admin/dictionaries/import/confirm", data={}).status_code == 503


@pytest.mark.parametrize("host,peer", [("127.0.0.1", "127.0.0.1"), ("localhost", "::1"), ("[::1]", "::1")])
def test_dev_admin_accepts_ipv4_and_ipv6_loopback(client, monkeypatch, host, peer):
    protected, _ = client
    monkeypatch.setenv("PROJECT2_DEV_ADMIN_OPEN", "1")
    monkeypatch.delenv("PROJECT2_ADMIN_TOKEN")
    # Host header exercises IPv6 parsing without TestClient's IPv6-netloc bug.
    with TestClient(protected.app, base_url="http://localhost", headers={"Host": host}, client=(peer, 50000)) as local:
        assert local.get("/admin").status_code == 200


@pytest.mark.parametrize("flag", [None, "0", "true"])
def test_dev_admin_disabled_retains_local_auth_and_csrf(client, monkeypatch, flag):
    protected, token = client
    if flag is not None:
        monkeypatch.setenv("PROJECT2_DEV_ADMIN_OPEN", flag)
    with TestClient(protected.app, base_url="http://localhost", client=("127.0.0.1", 50000)) as local:
        assert local.get("/admin").status_code == 403  # Original HTTPS requirement.
        monkeypatch.setenv("PROJECT2_ADMIN_ALLOW_HTTP", "1")
        assert local.get("/admin/dictionaries").status_code == 401
        login(local, token)
        assert local.post("/admin/dictionaries/new", data=dictionary_data()).status_code == 403


def test_dev_switch_retains_full_remote_token_and_csrf_protection(client, monkeypatch):
    remote, token = client
    monkeypatch.setenv("PROJECT2_DEV_ADMIN_OPEN", "1")
    assert remote.get("/admin/dictionaries").status_code == 401
    csrf = login(remote, token)
    assert remote.post("/admin/dictionaries/new", data=dictionary_data()).status_code == 403
    assert remote.post("/admin/dictionaries/new", data={**dictionary_data(), "csrf": csrf}).status_code == 200


def test_explicit_personal_edit_marks_complete_forms_verified_and_missing_forms_unknown(factory):
    with factory() as db:
        person = register_person(db, person_data())
        name = person.names[0]
        _, _, _, data = name_cases.personal_page(db, name.person_name_id)
        data.update({f"name_{case}": "Особистого" for case in CASES if case != "nominative"})
        data["name_verified"] = "on"
        name_cases.update_personal(db, name.person_name_id, data)
        row = db.scalar(select(PersonNameCases).where(PersonNameCases.kind == "name"))
        assert row.is_verified
        resolved = name_cases.resolve(db, name, "male")
        assert resolved["name"]["forms"]["genitive"] == "Особистого"
        assert resolved["surname"]["forms"]["genitive"] is None
        assert resolved["surname"]["forms"]["nominative"] == "Тестенко"
