"""Усі Excel-файли й SQLite-бази тут тимчасові, дані — штучні."""

from pathlib import Path
import subprocess
import sys

from openpyxl import Workbook
import pytest
from sqlalchemy import URL, create_engine, event, func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import sessionmaker

from app.database.base import Base
from app.models import Person, PersonName, PersonRank, Rank
from app.models.reference_data.ranks import ReferenceRanks
from app.services.seed_test_people import REQUIRED_COLUMNS, seed_test_people


@pytest.fixture
def isolated_db(tmp_path):
    path = tmp_path / "isolated.sqlite3"
    engine = create_engine(URL.create("sqlite", database=str(path)))

    @event.listens_for(engine, "connect")
    def enable_foreign_keys(connection, _):
        connection.execute("PRAGMA foreign_keys=ON")

    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False)
    with factory.begin() as db:
        db.add_all(Rank(**data) for data in ReferenceRanks.RANKS)
    yield factory, path
    engine.dispose()


def person_row(**overrides):
    data = dict(
        surname="Тестенко", name="Тестослав", patronymic="Тестович",
        tax_number="0123456789", birth_day=29, birth_month=2,
        birth_year=2000, rank="  СОЛДАТ  ", sex="male",
    )
    data.update(overrides)
    return [data[key] for key in REQUIRED_COLUMNS]


def workbook(tmp_path, rows, *, headers=REQUIRED_COLUMNS, sheet="persons"):
    book = Workbook()
    page = book.active
    page.title = sheet
    page.append(list(headers))
    page.append(["Підпис"] * len(headers))
    for row in rows:
        page.append(row)
    path = tmp_path / "people.xlsx"
    book.save(path)
    book.close()
    return path


def counts(factory):
    with factory() as db:
        return tuple(db.scalar(select(func.count()).select_from(model))
                     for model in (Person, PersonName, PersonRank))


def test_import_links_values_and_repeat_does_not_overwrite(isolated_db, tmp_path):
    factory, _ = isolated_db
    path = workbook(tmp_path, [person_row()])
    report = seed_test_people(factory, path)
    assert (report.found, report.created, report.skipped, report.issues) == (1, 1, 0, [])
    with factory.begin() as db:
        person = db.scalar(select(Person))
        assert person.tax_number == "0123456789"
        assert (person.birth_day, person.birth_month, person.birth_year) == (29, 2, 2000)
        assert len(person.names) == len(person.ranks) == 1
        name = person.names[0]
        assert (name.surname, name.name, name.patronymic) == ("Тестенко", "Тестослав", "Тестович")
        assert name.is_active and person.ranks[0].is_active
        assert person.ranks[0].rank.code == "SOLDIER"
        assert person.ranks[0].rank_order_number is None
        assert person.ranks[0].rank_order_date is None
        assert (person.ranks[0].rank_order_day, person.ranks[0].rank_order_month, person.ranks[0].rank_order_year) == (None, None, None)
        assert person.positions == person.service_history == []
        # Збережені дані не повинні замінюватися навіть значеннями з Excel.
        person.birth_year = 2004
        name.surname = "Зміненко"
        person.ranks[0].rank = db.scalar(select(Rank).where(Rank.code == "SENIOR_SOLDIER"))
        person.ranks[0].rank_order_year = 1995
        person.ranks[0].rank_order_month = 3
    again = seed_test_people(factory, path)
    assert (again.found, again.created, again.existing, again.skipped, again.issues) == (1, 0, 1, 1, [])
    assert counts(factory) == (1, 1, 1)
    with factory() as db:
        person = db.scalar(select(Person))
        assert person.birth_year == 2004
        assert person.names[0].surname == "Зміненко"
        assert person.ranks[0].rank.code == "SENIOR_SOLDIER"
        assert (person.ranks[0].rank_order_day, person.ranks[0].rank_order_month, person.ranks[0].rank_order_year) == (None, 3, 1995)
        assert person.ranks[0].rank_order_date is None


@pytest.mark.parametrize("day,month,year", [
    (None, 2, 2000), (29, 2, None), (None, None, 2000),
    (None, None, None), (29, None, 2000),
])
def test_preserves_partial_dates(isolated_db, tmp_path, day, month, year):
    factory, _ = isolated_db
    path = workbook(tmp_path, [person_row(birth_day=day, birth_month=month, birth_year=year)])
    report = seed_test_people(factory, path)
    assert report.created == 1 and not report.issues
    with factory() as db:
        person = db.scalar(select(Person))
        assert (person.birth_day, person.birth_month, person.birth_year) == (day, month, year)


def test_empty_rows_extra_column_and_duplicate_identifier(isolated_db, tmp_path):
    factory, _ = isolated_db
    path = workbook(tmp_path, [
        person_row() + [1], [None] * 9, [" "] * 9,
        person_row() + [2], person_row(tax_number="1234567890", rank="soldier") + [3],
    ], headers=(*REQUIRED_COLUMNS, "number"))
    report = seed_test_people(factory, path)
    assert (report.found, report.created, report.existing, report.skipped, report.issues) == (3, 2, 1, 1, [])
    assert counts(factory) == (2, 2, 2)


@pytest.mark.parametrize("value,number_format,expected", [
    (1234567890, "General", "1234567890"),
    (123456789, "0000000000", "0123456789"),
    ("0012345678", "General", "0012345678"),
])
def test_excel_tax_representation(isolated_db, tmp_path, value, number_format, expected):
    factory, _ = isolated_db
    path = workbook(tmp_path, [person_row(tax_number=value)])
    from openpyxl import load_workbook
    book = load_workbook(path)
    book["persons"].cell(3, REQUIRED_COLUMNS.index("tax_number") + 1).number_format = number_format
    book.save(path)
    book.close()
    assert seed_test_people(factory, path).created == 1
    with factory() as db:
        assert db.scalar(select(Person.tax_number)) == expected


@pytest.mark.parametrize("changes,message", [
    ({"tax_number": None}, "без РНОКПП"),
    ({"tax_number": 123456789}, "10 цифр"),
    ({"tax_number": "012345678X"}, "10 цифр"),
    ({"tax_number": 1234567890.5}, "ціле число"),
    ({"tax_number": True}, "логічне значення"),
    ({"tax_number": '=TEXT(123456789,"0000000000")'}, "формула"),
    ({"surname": None}, "ПІБ"),
    ({"name": "О1тест"}, "ПІБ"),
    ({"sex": None}, "стать"),
    ({"sex": "unknown"}, "стать"),
    ({"birth_day": 31, "birth_month": 4}, "Дата народження"),
    ({"birth_year": 2001}, "Дата народження"),
    ({"birth_month": 13}, "Дата народження"),
    ({"birth_day": 1.5}, "birth_day"),
    ({"rank": "Неіснуюче звання"}, "відсутнє в довіднику"),
    ({"rank": None}, "відсутнє в довіднику"),
])
def test_invalid_row_reports_number_and_continues(isolated_db, tmp_path, changes, message):
    factory, _ = isolated_db
    path = workbook(tmp_path, [person_row(**changes), [None] * 8,
                              person_row(tax_number="1234567890")])
    report = seed_test_people(factory, path)
    assert (report.found, report.created, report.existing, report.skipped) == (2, 1, 0, 1)
    assert len(report.issues) == 1
    assert report.issues[0].row_number == 3
    assert message in report.issues[0].message
    assert "0123456789" not in report.issues[0].message
    assert "Тестенко" not in report.issues[0].message
    assert counts(factory) == (1, 1, 1)


def test_unknown_rank_is_not_created(isolated_db, tmp_path):
    factory, _ = isolated_db
    path = workbook(tmp_path, [person_row(rank="Неіснуюче звання")])
    assert len(seed_test_people(factory, path).issues) == 1
    assert counts(factory) == (0, 0, 0)
    with factory() as db:
        assert db.scalar(select(func.count()).select_from(Rank)) == len(ReferenceRanks.RANKS)


def test_ambiguous_rank_requires_code(isolated_db, tmp_path):
    factory, _ = isolated_db
    with factory.begin() as db:
        source = next(row for row in ReferenceRanks.RANKS if row["code"] == "SOLDIER")
        db.add(Rank(**{**source, "code": "SOLDIER_DUPLICATE"}))
    path = workbook(tmp_path, [person_row()])
    assert "неоднозначне" in seed_test_people(factory, path).issues[0].message
    assert counts(factory) == (0, 0, 0)
    path = workbook(tmp_path, [person_row(rank="SOLDIER")])
    assert seed_test_people(factory, path).created == 1


def test_error_after_related_insert_rolls_back_entire_person(isolated_db, tmp_path):
    factory, _ = isolated_db
    path = workbook(tmp_path, [person_row(), person_row(tax_number="1234567890")])
    attempts = []

    def fail_first_rank(*_):
        attempts.append(True)
        if len(attempts) == 1:
            raise IntegrityError("private SQL", {"tax_number": "0123456789"}, Exception("private data"))

    event.listen(PersonRank, "after_insert", fail_first_rank)
    try:
        report = seed_test_people(factory, path)
    finally:
        event.remove(PersonRank, "after_insert", fail_first_rank)
    assert (report.created, report.skipped) == (1, 1)
    assert report.issues[0].row_number == 3
    assert "цілісності" in report.issues[0].message
    assert "private" not in report.issues[0].message
    assert "0123456789" not in report.issues[0].message
    assert counts(factory) == (1, 1, 1)
    with factory() as db:
        assert db.scalar(select(Person.tax_number)) == "1234567890"
    # Після усунення помилки перший рядок можна безпечно повторити.
    again = seed_test_people(factory, path)
    assert (again.created, again.existing, again.issues) == (1, 1, [])
    assert counts(factory) == (2, 2, 2)


@pytest.mark.parametrize("sheet,headers,message", [
    ("positions", REQUIRED_COLUMNS, "немає аркуша persons"),
    ("persons", REQUIRED_COLUMNS[:-1], "відсутні колонки: sex"),
    ("persons", (*REQUIRED_COLUMNS, "rank"), "повторюються колонки: rank"),
    ("persons", (), "відсутні колонки"),
])
def test_structural_errors_before_opening_session(tmp_path, sheet, headers, message):
    path = workbook(tmp_path, [], sheet=sheet, headers=headers)
    # None замість фабрики: жодна сесія не повинна відкриватися.
    with pytest.raises(ValueError, match=message):
        seed_test_people(None, path)


def test_empty_sheet_has_zero_counts(isolated_db, tmp_path):
    factory, _ = isolated_db
    report = seed_test_people(factory, workbook(tmp_path, []))
    assert (report.found, report.created, report.skipped, report.issues) == (0, 0, 0, [])
    assert counts(factory) == (0, 0, 0)


def test_invalid_source_for_existing_person_is_reported_without_changes(isolated_db, tmp_path):
    factory, _ = isolated_db
    path = workbook(tmp_path, [person_row()])
    assert seed_test_people(factory, path).created == 1
    path = workbook(tmp_path, [person_row(rank="Неіснуюче звання")])
    report = seed_test_people(factory, path)
    assert report.created == 0 and report.skipped == 1
    assert report.issues[0].row_number == 3
    assert counts(factory) == (1, 1, 1)


def test_corrupt_workbook_is_reported_before_opening_session(tmp_path):
    path = tmp_path / "broken.xlsx"
    path.write_text("not an Excel workbook")
    with pytest.raises(ValueError, match="коректний файл .xlsx"):
        seed_test_people(None, path)


def run_cli(excel_path, database_path):
    return subprocess.run(
        [sys.executable, "-m", "app.services.seed_test_people",
         "--excel", str(excel_path), "--database", str(database_path)],
        cwd=Path(__file__).resolve().parents[1], capture_output=True, text=True,
    )


def test_module_cli_repeat_and_exit_codes(isolated_db, tmp_path):
    factory, path = isolated_db
    excel = workbook(tmp_path, [person_row()])
    first = run_cli(excel, path)
    assert first.returncode == 0, first.stderr
    assert "створено: 1" in first.stdout and "помилок: 0" in first.stdout
    second = run_cli(excel, path)
    assert second.returncode == 0, second.stderr
    assert "створено: 0" in second.stdout and "вже існують: 1" in second.stdout
    invalid = workbook(tmp_path, [person_row(tax_number="bad")])
    failed_row = run_cli(invalid, path)
    assert failed_row.returncode == 1
    assert "Рядок 3: tax_number:" in failed_row.stdout
    assert counts(factory) == (1, 1, 1)
    missing = tmp_path / "must_not_be_created.sqlite3"
    failed_file = run_cli(excel, missing)
    assert failed_file.returncode == 2 and "файл SQLite не існує" in failed_file.stderr
    assert not missing.exists()
    no_sheet = workbook(tmp_path, [], sheet="positions")
    failed_sheet = run_cli(no_sheet, path)
    assert failed_sheet.returncode == 2 and "немає аркуша persons" in failed_sheet.stderr
    assert counts(factory) == (1, 1, 1)
