"""Existing positions worksheet mapping, preview without DB writes, atomic import."""
from datetime import date
from openpyxl import load_workbook
from zipfile import ZipFile, BadZipFile
from app.validators.position_validator import POSITION_ATTRIBUTES, validate_position
from app.repositories.staff_repository import StaffRepository
from app.models.staff import StaffStructure, StaffVersion
from app.models.reference_book import Position
from app.services.staff_service import period, version_conflict, write_unit, checked


def read_positions(source):
    errors=[]; rows=[]; seen=set()
    try:
        with ZipFile(source) as archive:
            if sum(item.file_size for item in archive.infolist())>50*1024*1024:
                raise ValueError("Розпакований Excel перевищує 50 МБ.")
        if hasattr(source,"seek"): source.seek(0)
        book=load_workbook(source,read_only=True,data_only=True)
    except (BadZipFile, OSError, KeyError) as error:
        raise ValueError("Не вдалося прочитати файл .xlsx.") from error
    try:
        if "positions" not in book.sheetnames: raise ValueError("Аркуш positions відсутній.")
        sheet=book["positions"]
        header=[str(value or "").strip().lower().replace(" ","_").replace("-","_") for value in next(sheet.iter_rows(values_only=True),())]
        if len([value for value in header if value])!=len(set(value for value in header if value)):
            raise ValueError("Повторні назви колонок Excel.")
        missing=set(POSITION_ATTRIBUTES)-set(header)
        if missing: raise ValueError("Відсутні колонки: "+", ".join(sorted(missing)))
        # Row 2 is explanatory in the repository's established workbook format.
        for number,values in enumerate(sheet.iter_rows(min_row=3,values_only=True),3):
            if not any(value is not None and str(value).strip() for value in values): continue
            try:
                row=validate_position(dict(zip(header,values)))
                if row["position_number"] in seen: raise ValueError("Повторний номер посади.")
                seen.add(row["position_number"]); rows.append(row)
            except ValueError as error: errors.append(f"Рядок {number}: {error}")
        if not rows: errors.append("Немає коректних посад для імпорту.")
    finally: book.close()
    return {"positions":rows,"errors":errors,"units":len({row["unit_short"] for row in rows}),
            "warnings":["Підрозділи визначено за unit_short; окремих ієрархічних кодів підрозділів у цьому форматі немає."]}


def preview_import(db, source, structure_id, name, start, end, current):
    period(start,end)
    if not name.strip() or len(name.strip())>150: raise ValueError("Назва версії: 1–150 символів.")
    if current:
        if end is not None: raise ValueError("Для поточної версії завершення має бути порожнім.")
        from app.services.staff_service import today
        if start>today(): raise ValueError("Поточна версія не може починатися в майбутньому.")
    elif end is None:
        raise ValueError("Для архівної версії вкажіть точне завершення.")
    if structure_id is not None: checked(db,StaffStructure,structure_id)
    result=read_positions(source)
    if structure_id is not None:
        try: version_conflict(db,structure_id,start,end)
        except ValueError as error: result["errors"].append(str(error))
    result.update(structure_id=structure_id,name=name.strip(),start=start.isoformat(),end=end.isoformat() if end else None)
    return result


def import_version(db, prepared, import_key):
    if prepared["errors"]: raise ValueError("Спочатку виправте критичні помилки імпорту.")
    with write_unit(db):
        return insert_prepared(db, prepared, import_key)


def insert_prepared(db, prepared, import_key):
    start=date.fromisoformat(prepared["start"]);end=date.fromisoformat(prepared["end"]) if prepared["end"] else None
    if StaffRepository.version_by_import_key(db, import_key):
        raise ValueError("Це завантаження вже підтверджено.")
    structure_id=prepared["structure_id"]
    if structure_id is None:
        structure=StaffRepository.structure_by_code(db, "main")
        if structure is None:
            structure=StaffStructure(code="main",name="Основний штат");StaffRepository.add(db, structure);db.flush()
        structure_id=structure.id
    checked(db,StaffStructure,structure_id)
    version_conflict(db,structure_id,start,end)
    if StaffRepository.version_by_name(db, structure_id, prepared["name"]):
        raise ValueError("Версія з такою назвою вже існує.")
    version=StaffVersion(staff_structure_id=structure_id,name=prepared["name"],valid_from=start,valid_to=end,import_key=import_key)
    StaffRepository.add(db, version);db.flush()
    for row in prepared["positions"]: StaffRepository.add(db, Position(staff_version_id=version.id,**validate_position(row)))
    db.flush(); key=version.id
    return key
