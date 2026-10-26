"""Окремий, ідемпотентний імпорт тестових людей з аркуша persons."""

import argparse
from dataclasses import dataclass, field
from pathlib import Path
import sys
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import URL, create_engine
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from sqlalchemy.orm import sessionmaker

from app.config import EXCEL_PATH
from app.repositories.person_repository import PersonRepository
from app.services.service_person_group_data_validation import DataGroupValidate
from app.validators.person_validator import PersonValidator
from app.database.schema_check import require_current_schema
from app.date_fields import DATE_FIELDS
from app.utils.date_utils import date_values_from_input


REQUIRED_COLUMNS = (
    "surname", "name", "patronymic", "tax_number",
    "birth_day", "birth_month", "birth_year", "rank",
)


@dataclass(frozen=True)
class RowIssue:
    row_number: int
    message: str


@dataclass
class SeedReport:
    found: int = 0
    created: int = 0
    existing: int = 0
    issues: list[RowIssue] = field(default_factory=list)

    @property
    def skipped(self) -> int:
        """Усі нестворені рядки: вже наявні люди та рядки з помилками."""
        return self.existing + len(self.issues)


def _is_empty(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _cell_text(cell, column: str) -> str | None:
    if _is_empty(cell.value):
        return None
    if cell.data_type in ("f", "e"):
        raise ValueError(f"{column}: формула або помилка Excel; потрібне значення.")
    value = cell.value
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, bool):
        raise ValueError(f"{column}: логічне значення не підтримується.")
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not value.is_integer():
            raise ValueError(f"{column}: потрібне ціле число або текст.")
        text = str(int(value))
        # Відновлюємо лише нулі, явно задані форматом Excel, без здогадок.
        if column == "tax_number" and cell.number_format == "0000000000":
            text = text.zfill(10)
        return text
    raise ValueError(f"{column}: непідтримуваний тип значення Excel.")


def _read_rows(excel_path: Path) -> list[tuple[int, dict]]:
    try:
        workbook = load_workbook(excel_path, read_only=True, data_only=False)
    except (BadZipFile, InvalidFileException) as error:
        raise ValueError("Не вдалося прочитати Excel; потрібен коректний файл .xlsx.") from error
    try:
        if "persons" not in workbook.sheetnames:
            raise ValueError("У файлі Excel немає аркуша persons.")
        sheet = workbook["persons"]
        rows = sheet.iter_rows()
        header = next(rows, ())
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in header]
        missing = [key for key in REQUIRED_COLUMNS if key not in headers]
        if missing:
            raise ValueError("Аркуш persons: відсутні колонки: " + ", ".join(missing))
        duplicates = [key for key in REQUIRED_COLUMNS if headers.count(key) > 1]
        if duplicates:
            raise ValueError("Аркуш persons: повторюються колонки: " + ", ".join(duplicates))
        indexes = {key: headers.index(key) for key in REQUIRED_COLUMNS}
        return [
            (number, {key: row[index] for key, index in indexes.items()})
            for number, row in enumerate(rows, 2)
            if any(not _is_empty(cell.value) for cell in row)
        ]
    finally:
        workbook.close()


def _rank_key(value: str) -> str:
    return " ".join(value.split()).casefold()


def _create_row(db, data: dict) -> bool:
    try:
        tax_number = PersonValidator.validate_tax_number(data["tax_number"])
    except ValueError as error:
        raise ValueError("tax_number: РНОКПП має містити рівно 10 цифр.") from error
    if not tax_number:
        # У моделі РНОКПП nullable, але ПІБ/дата не є унікальними ключами.
        raise ValueError("tax_number: без РНОКПП неможливо надійно уникнути дублікатів; рядок пропущено.")
    try:
        person_data = DataGroupValidate.validate_person_base_data(data)
    except ValueError as error:
        raise ValueError(f"Дата народження: {error}") from error
    try:
        name_data = DataGroupValidate.validate_name_data(data)
        if not name_data["surname"] or not name_data["name"]:
            raise ValueError("Required name")
    except ValueError as error:
        raise ValueError("ПІБ: surname і name обов'язкові; непорожні частини — лише літери, 3–25 символів.") from error

    rank_key = _rank_key(data["rank"] or "")
    if not rank_key:
        raise ValueError("rank: звання відсутнє в довіднику; потрібен code або називний відмінок.")
    matches = [
        rank for rank in PersonRepository.ranks(db)
        if rank_key in (_rank_key(rank.code), _rank_key(rank.nominative))
    ]
    if not matches:
        raise ValueError("rank: звання відсутнє в довіднику; потрібен code або називний відмінок.")
    if len(matches) != 1:
        raise ValueError("rank: неоднозначне звання в довіднику; вкажіть унікальний code.")
    if PersonRepository.get_by_tax_number(db, tax_number) is not None:
        return False

    person = PersonRepository.add_person(db, person_data)
    PersonRepository.add_name(db, person.person_id, name_data)
    # The persons sheet has no order date: keep all of its components unknown.
    rank_date = date_values_from_input(DATE_FIELDS["rank.rank_order_date"], "rank_order", {})
    PersonRepository.add_rank(db, person.person_id, dict(rank_id=matches[0].rank_id, is_active=True, **rank_date))
    db.flush()
    return True


def seed_test_people(session_factory, excel_path: Path = EXCEL_PATH) -> SeedReport:
    """Кожен рядок має власну транзакцію; наявні люди ніколи не оновлюються.

    session_factory має створювати нові SQLAlchemy Session. Структурні помилки
    Excel перевіряються до відкриття сесій. Схема БД та звання не створюються.
    """
    rows = _read_rows(Path(excel_path))
    report = SeedReport(found=len(rows))
    for number, cells in rows:
        try:
            data = {key: _cell_text(cell, key) for key, cell in cells.items()}
            with session_factory.begin() as db:
                created = _create_row(db, data)
            # Лічильник змінюється лише після успішного commit.
            if created:
                report.created += 1
            else:
                report.existing += 1
        except ValueError as error:
            report.issues.append(RowIssue(number, str(error)))
        except IntegrityError:
            # Не виводимо SQL/параметри запиту з персональними даними.
            report.issues.append(RowIssue(number, "Обмеження цілісності БД; усі зміни рядка скасовано."))
    return report


def main(argv=None) -> int:
    from app.database.db import DB_PATH

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--excel", type=Path, default=EXCEL_PATH, help="Шлях до reference_data.xlsx")
    parser.add_argument("--database", type=Path, default=DB_PATH, help="Шлях до вже підготовленої SQLite-бази")
    args = parser.parse_args(argv)
    if not args.database.is_file():
        print("Помилка: файл SQLite не існує; спочатку підготуйте схему та довідник звань.", file=sys.stderr)
        return 2

    engine = create_engine(URL.create("sqlite", database=str(args.database.resolve())))
    try:
        require_current_schema(engine)
        report = seed_test_people(sessionmaker(bind=engine, autoflush=False), args.excel)
    except (OSError, ValueError) as error:
        print(f"Помилка: {error}", file=sys.stderr)
        return 2
    except SQLAlchemyError:
        print("Помилка доступу до БД або несумісна схема; перевірте файл та наявність поточних таблиць.", file=sys.stderr)
        return 2
    finally:
        engine.dispose()

    print(f"Знайдено: {report.found}; створено: {report.created}; пропущено: {report.skipped} "
          f"(вже існують: {report.existing}; помилок: {len(report.issues)}).")
    for issue in report.issues:
        print(f"Рядок {issue.row_number}: {issue.message}")
    return 1 if report.issues else 0


if __name__ == "__main__":
    raise SystemExit(main())
