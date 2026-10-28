"""Bounded Excel preview and atomic, revision-checked directory import."""
import hashlib
import json
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy.exc import SQLAlchemyError
from app.database.transactions import service_transaction
from app.models.service_basis import TccOffice, TccImport
from app.repositories.service_basis_repository import TccRepository as Repo
from app.services.tcc_directory import values, validate_parent
from app.services.service_bases import reference_id

HEADERS = ('id', 'official_name', 'type', 'region', 'settlement', 'parent_tcc_id', 'is_active')
LABELS = ('ID існуючого запису (або порожньо)', 'Офіційна назва', 'Тип (системний ключ)', 'Область', 'Населений пункт', 'ID батьківського ТЦК', 'Активний (1 або 0)')


def empty_template():
    book = Workbook(); sheet = book.active; sheet.title = 'tcc_offices'
    sheet.append(HEADERS); sheet.append(LABELS); sheet.freeze_panes = 'A3'
    for cell in sheet[1]:
        sheet.column_dimensions[cell.column_letter].width = 32
    output = BytesIO(); book.save(output); book.close()
    return output.getvalue()


def integer(value, label):
    return reference_id(value, label)


def read_workbook(content):
    if len(content) > 5 * 1024 * 1024:
        raise ValueError('Файл перевищує 5 МБ.')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError('Розпакований файл перевищує 50 МБ.')
        book = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, OSError, KeyError, ValueError, InvalidFileException, ParseError) as error:
        raise ValueError('Неможливо прочитати Excel; перевірте формат і розмір .xlsx.') from error
    rows, errors, seen, ids = [], [], set(), set()
    try:
        if 'tcc_offices' not in book.sheetnames:
            raise ValueError('Немає аркуша tcc_offices.')
        sheet = book['tcc_offices']
        if (sheet.max_row or 0) > 10002:
            raise ValueError('Допускається до 10000 рядків даних.')
        header = [str(c.value or '').strip() for c in next(sheet.iter_rows(), ())]
        while header and not header[-1]:
            header.pop()
        if tuple(header) != HEADERS:
            raise ValueError('Рядок 1: потрібні заголовки ' + ', '.join(HEADERS))
        for number, cells in enumerate(sheet.iter_rows(min_row=3), 3):
            if not any(c.value is not None and str(c.value).strip() for c in cells):
                continue
            if len(rows) + len(errors) >= 10000:
                raise ValueError('Допускається до 10000 заповнених рядків.')
            try:
                if any(c.data_type == 'f' for c in cells):
                    raise ValueError('Формули заборонені.')
                if any(c.value is not None for c in cells[len(HEADERS):]):
                    raise ValueError('Є дані за межами заголовків.')
                data = dict(zip(HEADERS, [c.value for c in cells]))
                for field in ('official_name', 'type', 'region', 'settlement'):
                    if data.get(field) is not None and not isinstance(data[field], str):
                        raise ValueError(f'{field}: потрібне текстове значення.')
                key = integer(data.get('id'), 'id')
                data['parent_tcc_id'] = integer(data.get('parent_tcc_id'), 'parent_tcc_id')
                if isinstance(data.get('is_active'), bool) or data.get('is_active') not in (0, 1, '0', '1'):
                    raise ValueError('is_active: вкажіть 1 або 0.')
                prepared = values({**data, 'is_active': str(data['is_active'])})
                identity = tuple(prepared[f] for f in ('normalized_name', 'normalized_region', 'normalized_settlement'))
                if identity in seen or (key is not None and key in ids):
                    raise ValueError('Дублікат ТЦК або ID у файлі.')
                seen.add(identity); ids.add(key)
                rows.append({'row': number, 'key': key, 'values': prepared})
            except ValueError as error:
                errors.append(f'tcc_offices, рядок {number}: {error}')
    finally:
        book.close()
    if not rows and not errors:
        errors.append('Немає заповнених рядків.')
    return rows, errors


def preview(db, content, mode):
    if mode not in ('add', 'update'):
        raise ValueError('Оберіть режим додавання або оновлення.')
    rows, errors = read_workbook(content)
    counts = {key: 0 for key in ('new', 'updated', 'unchanged', 'skipped', 'errors')}
    proposed = {}
    seen_ids = set()
    for row in rows:
        try:
            by_identity = Repo.find(db, row['values'])
            old = Repo.get(db, row['key']) if row['key'] else by_identity
            if row['key'] and old is None:
                raise ValueError('Вказаний ID не існує; для нового ТЦК залиште id порожнім.')
            if by_identity is not None and old is not None and by_identity.id != old.id:
                raise ValueError('Назва та місце вже належать іншому ID.')
            if old and old.id in seen_ids:
                raise ValueError('Кілька рядків змінюють той самий ТЦК.')
            if old: seen_ids.add(old.id)
            row['key'] = old.id if old else None
            row['expected'] = old.revision if old else None
            status = 'new' if old is None else ('unchanged' if all(getattr(old, k) == v for k, v in row['values'].items()) else ('updated' if mode == 'update' else 'skipped'))
            row['status'] = status; counts[status] += 1
            if old and status == 'updated': proposed[old.id] = row['values']['parent_tcc_id']
        except ValueError as error:
            errors.append(f'tcc_offices, рядок {row["row"]}: {error}')
    for row in rows:
        if row.get('status') not in ('new', 'updated'):
            continue
        try:
            validate_parent(db, row['values']['parent_tcc_id'], row['key'], proposed)
        except ValueError as error:
            errors.append(f'tcc_offices, рядок {row["row"]}: {error}')
    counts['errors'] = len(errors)
    digest = hashlib.sha256(json.dumps(rows, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    return {'resource': 'tcc', 'rows': rows, 'errors': errors, 'counts': counts, 'mode': mode, 'digest': digest}


def confirm(db, prepared, import_id, *, reviewed=False, confirm_updates=False):
    if prepared.get('resource') != 'tcc' or prepared['errors'] or not prepared['rows']:
        raise ValueError('Потрібен коректний попередній перегляд довідника ТЦК.')
    if not reviewed or (prepared['counts']['updated'] and not confirm_updates):
        raise ValueError('Підтвердьте перевірку та окремо оновлення існуючих установ.')
    if not (prepared['counts']['new'] or prepared['counts']['updated']):
        raise ValueError('Немає змін для імпорту.')
    try:
        with service_transaction(db, immediate=True):
            if Repo.imported(db, import_id):
                raise ValueError('Імпорт уже виконано; повторне підтвердження заборонене.')
            proposed = {row['key']: row['values']['parent_tcc_id'] for row in prepared['rows'] if row['key'] and row['status'] == 'updated'}
            for row in prepared['rows']:
                old = Repo.get(db, row['key']) if row['key'] else Repo.find(db, row['values'])
                if (old.revision if old else None) != row['expected']:
                    raise ValueError('Довідник змінився після перевірки. Повторіть preview.')
                if row['status'] not in ('new', 'updated'):
                    continue
                validate_parent(db, row['values']['parent_tcc_id'], row['key'], proposed)
                record = old or Repo.add(db, TccOffice())
                for key, value in row['values'].items(): setattr(record, key, value)
                db.flush()
            Repo.add(db, TccImport(id=import_id, digest=prepared['digest'], mode=prepared['mode']))
        return prepared['counts']
    except SQLAlchemyError:
        raise ValueError('Конфлікт імпорту: усі зміни скасовано. Повторіть перевірку.') from None
