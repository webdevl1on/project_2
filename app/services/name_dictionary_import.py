"""Three explicit Excel sheets, reviewed preview, one atomic confirmation."""
import hashlib
import json
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from xml.etree.ElementTree import ParseError
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm.exc import StaleDataError
from app.database.transactions import service_transaction
from app.models.name_cases import NameDictionary, NameDictionaryImport
from app.name_cases import SHEETS, IMPORT_HEADERS
from app.repositories.name_case_repository import NameCaseRepository as Repository
from app.services.name_cases import dictionary_values


def empty_template():
    book = Workbook()
    book.remove(book.active)
    for title in SHEETS:
        sheet = book.create_sheet(title)
        sheet.append(IMPORT_HEADERS)
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.comment = Comment("Обов'язкове текстове поле. sex: male або female. base_form дорівнює nominative.", "project_2")
            sheet.column_dimensions[cell.column_letter].width = 24
        validation = DataValidation(type="list", formula1='"male,female"', allow_blank=False)
        validation.errorTitle = "Стать не визначається автоматично"
        validation.error = "Вкажіть male або female."
        validation.showErrorMessage = True
        sheet.add_data_validation(validation)
        validation.add("B2:B10001")
    target = BytesIO()
    book.save(target)
    book.close()
    return target.getvalue()


def read_workbook(content):
    if len(content) > 5 * 1024 * 1024:
        raise ValueError("Файл перевищує 5 МБ.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("Розпакований Excel перевищує 50 МБ.")
        book = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, OSError, KeyError, ValueError, InvalidFileException, ParseError) as error:
        raise ValueError("Не вдалося прочитати .xlsx; перевірте формат і розмір файлу.") from error
    rows = []; errors = []; seen = set()
    try:
        missing = set(SHEETS) - set(book.sheetnames)
        if missing:
            raise ValueError("Відсутні аркуші: " + ", ".join(sorted(missing)))
        for sheet_name, kind in SHEETS.items():
            sheet = book[sheet_name]
            if sheet.max_row and sheet.max_row > 10001:
                raise ValueError(f"{sheet_name}: за один імпорт допускається до 10000 рядків даних.")
            header = [str(cell.value or "").strip() for cell in next(sheet.iter_rows(), ())]
            while header and not header[-1]:
                header.pop()
            if tuple(header) != IMPORT_HEADERS:
                errors.append(f"{sheet_name}, рядок 1: потрібні заголовки у порядку: {', '.join(IMPORT_HEADERS)}.")
                continue
            for number, cells in enumerate(sheet.iter_rows(min_row=2), 2):
                if not any(cell.value is not None and str(cell.value).strip() for cell in cells):
                    continue
                if len(rows) + len(errors) >= 10000:
                    raise ValueError("За один імпорт допускається до 10000 заповнених рядків.")
                try:
                    if any(cell.data_type == "f" for cell in cells):
                        raise ValueError("Формули не підтримуються; потрібні явні текстові значення.")
                    if any(cell.value is not None and str(cell.value).strip() for cell in cells[len(IMPORT_HEADERS):]):
                        raise ValueError("Є дані за межами визначених колонок.")
                    values = dictionary_values({"kind": kind, **dict(zip(IMPORT_HEADERS, (cell.value for cell in cells)))}, verified=True)
                    key = (kind, values["sex"], values["normalized_key"])
                    if key in seen:
                        raise ValueError("Повтор слова з тією самою статтю в одному словнику.")
                    seen.add(key)
                    rows.append({"sheet": sheet_name, "row": number, "values": values})
                except ValueError as error:
                    errors.append(f"{sheet_name}, рядок {number}: {error}")
    finally:
        book.close()
    if not rows and not errors:
        errors.append("Файл не містить заповнених рядків.")
    return rows, errors


def preview(db, content, mode):
    if mode not in ("add", "update"):
        raise ValueError("Оберіть режим додавання або оновлення.")
    rows, errors = read_workbook(content)
    counts = {"new": 0, "updated": 0, "unchanged": 0, "skipped": 0, "errors": len(errors)}
    for row in rows:
        values = row["values"]
        old = Repository.find(db, values["kind"], values["sex"], values["normalized_key"])
        row["expected"] = old.revision if old else None
        if old is None:
            status = "new"
        elif all(getattr(old, key) == value for key, value in values.items()):
            status = "unchanged"
        else:
            status = "updated" if mode == "update" else "skipped"
        row["status"] = status
        counts[status] += 1
    digest = hashlib.sha256(json.dumps([row["values"] for row in rows], ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    return {"mode": mode, "rows": rows, "errors": errors, "counts": counts, "digest": digest}


def confirm(db, prepared, import_id, *, confirm_updates=False, reviewed=False):
    if prepared["errors"] or not prepared["rows"]:
        raise ValueError("Спочатку виправте помилки файлу.")
    if not reviewed:
        raise ValueError("Підтвердьте перевірку довідкових форм.")
    if prepared["counts"]["updated"] and not confirm_updates:
        raise ValueError("Окремо підтвердьте оновлення існуючих словникових записів.")
    if not (prepared["counts"]["new"] or prepared["counts"]["updated"]):
        raise ValueError("Немає змін для імпорту.")
    try:
        with service_transaction(db, immediate=True):
            if Repository.imported(db, import_id):
                raise ValueError("Цей імпорт уже підтверджено. Повторний запис заборонено.")
            for row in prepared["rows"]:
                values = row["values"]
                old = Repository.find(db, values["kind"], values["sex"], values["normalized_key"])
                if (old.revision if old else None) != row["expected"]:
                    raise ValueError("Словник змінився після перегляду. Повторіть перевірку файлу.")
                if row["status"] not in ("new", "updated"):
                    continue
                record = old or Repository.add(db, NameDictionary())
                for key, value in values.items():
                    setattr(record, key, value)
                db.flush()  # Detect every constraint failure before the one final commit.
            Repository.add(db, NameDictionaryImport(id=import_id, digest=prepared["digest"], mode=prepared["mode"]))
        return prepared["counts"]
    except (IntegrityError, StaleDataError) as error:
        raise ValueError("Конфлікт імпорту; усі зміни скасовано. Повторіть попередній перегляд.") from error
