"""Optional synthetic assignments; dates must be explicit, never inferred."""
from datetime import date
from openpyxl import load_workbook
from app.services.staff_service import StaffService, period

HEADERS=('tax_number','position_number','date_start','date_end')


def read_assignments(excel, config, position_numbers):
    book=load_workbook(excel,read_only=True,data_only=False)
    try:
        if 'assignments' not in book.sheetnames: return []
        sheet=book['assignments']
        if tuple(next(sheet.iter_rows(values_only=True),()))!=HEADERS: raise ValueError('assignments: неправильні заголовки.')
        rows=[];seen=set()
        for number,values in enumerate(sheet.iter_rows(min_row=3,values_only=True),3):
            if not any(v is not None for v in values): continue
            try:
                tax,position,start,end=values
                if not isinstance(tax,str) or len(tax)!=10 or not tax.isdigit(): raise ValueError('Потрібен текстовий тестовий РНОКПП.')
                position=int(position)
                if position not in position_numbers: raise ValueError('Невідома посада.')
                start=date.fromisoformat(start) if isinstance(start,str) else None
                end=date.fromisoformat(end) if isinstance(end,str) else None if end is None else False
                period(start,end)
                if start<config['valid_from'] or (config['valid_to'] and (end is None or end>config['valid_to'])): raise ValueError('Призначення поза межами версії.')
                record={'tax_number':tax,'position_number':position,'start':start,'end':end}
                key=(tax,position,start,end)
                if key in seen: raise ValueError('Повтор призначення.')
                seen.add(key);rows.append(record)
            except (ValueError,TypeError) as error: raise ValueError(f'assignments, рядок {number}: {error}') from None
        for i,row in enumerate(rows):
            for other in rows[:i]:
                if row['tax_number']==other['tax_number'] or row['position_number']==other['position_number']:
                    if max(row['start'],other['start'])<=min(row['end'] or date.max,other['end'] or date.max):
                        raise ValueError('assignments: конфлікт періодів.')
        return rows
    finally:book.close()


def seed_assignments(factory, rows, version_id):
    count=0
    for row in rows:
        with factory() as db:
            from app.repositories.seed_repository import assignment_targets
            person_id,position_id,existing=assignment_targets(db,row,version_id)
            if person_id is None or position_id is None: raise ValueError('assignments: невідоме посилання.')
            if existing is None:
                StaffService.create_assignment(db,person_id,position_id,row['start'],row['end']);count+=1
    return count
