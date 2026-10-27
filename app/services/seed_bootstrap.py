"""Preflight and ordered population of a staging database, never the published file."""
from datetime import date
from pathlib import Path
from time import monotonic
from openpyxl import load_workbook
from sqlalchemy import create_engine, URL
from sqlalchemy.orm import sessionmaker
from app.database.init_db import init_db
from app.database.schema_check import require_current_schema
from app.seeds.seed_structure import seed_all_structure
from app.services.seed_packages import preflight_names, preflight_geography, sha256, check_hashes
from app.services.seed_dictionaries import seed_name_package
from app.services.seed_geography import seed_geography
from app.services.seed_test_people import _read_rows, _cell_text, REQUIRED_COLUMNS, seed_test_people, _rank_key
from app.services.service_person_group_data_validation import DataGroupValidate
from app.validators.person_validator import PersonValidator
from app.validators.name_cases import sex
from app.models.reference_data.ranks import ReferenceRanks
from app.services.staff_import import read_positions
from app.services.staff_service import period, StaffService


def read_config(excel):
    book=load_workbook(excel,read_only=True,data_only=False)
    try:
        if 'seed_config' not in book.sheetnames: raise ValueError('Потрібен seed_config з явною датою початку штату.')
        sheet=book['seed_config']; headers=next(sheet.iter_rows(values_only=True),())
        if tuple(headers)!=('version_name','structure_code','valid_from','valid_to'): raise ValueError('seed_config: неправильні заголовки.')
        rows=[r for r in sheet.iter_rows(min_row=3,values_only=True) if any(v is not None for v in r)]
        if len(rows)!=1: raise ValueError('seed_config: потрібен один запис у рядку 3.')
        values=dict(zip(headers,rows[0]))
        for key in ('valid_from','valid_to'):
            raw=values[key]
            if raw is None and key=='valid_to': values[key]=None
            elif isinstance(raw,str):
                try: values[key]=date.fromisoformat(raw)
                except ValueError: raise ValueError(f'seed_config: {key} потребує повної дати.') from None
            elif type(raw) is date: values[key]=raw
            else: raise ValueError(f'seed_config: {key} потребує явно заданої повної дати.')
        period(values['valid_from'],values['valid_to'])
        if not all(isinstance(values[k],str) and values[k].strip() for k in ('version_name','structure_code')):
            raise ValueError('seed_config: назва й код обов’язкові.')
        if len(values['version_name'])>150: raise ValueError('seed_config: назва версії надто довга.')
        from app.services.staff_service import today
        if values['valid_to'] is None and values['valid_from']>today(): raise ValueError('seed_config: поточний штат не може починатися у майбутньому.')
        return values
    finally: book.close()


def preflight(excel, names, geography):
    excel=Path(excel); names=Path(names); geography=Path(geography)
    config=read_config(excel); positions=read_positions(excel)
    if positions['errors']: raise ValueError('positions: '+'; '.join(positions['errors']))
    ranks={_rank_key(r[k]) for r in ReferenceRanks.RANKS for k in ('code','nominative')}
    rows=_read_rows(excel); seen=set()
    for number,cells in rows:
        try:
            values={key:_cell_text(cell,key) for key,cell in cells.items()}
            tax=PersonValidator.validate_tax_number(values['tax_number'])
            if not tax or tax in seen: raise ValueError('Відсутній або повторний РНОКПП у тестових рядках.')
            seen.add(tax); sex(values['sex'])
            DataGroupValidate.validate_person_base_data(values)
            names_data=DataGroupValidate.validate_name_data(values)
            if not names_data['surname'] or not names_data['name']: raise ValueError('ПІБ обов’язковий.')
            if _rank_key(values['rank'] or '') not in ranks: raise ValueError('Звання відсутнє в довіднику.')
        except ValueError as error: raise ValueError(f'persons, рядок {number}: {error}') from None
    if not rows: raise ValueError('persons: немає тестових людей.')
    from app.services.seed_assignments import read_assignments
    assignments=read_assignments(excel,config,{p['position_number'] for p in positions['positions']})
    if any(row['tax_number'] not in seen for row in assignments): raise ValueError('assignments: невідома тестова людина.')
    name_package=preflight_names(names); geo_package=preflight_geography(geography)
    return {'config':config,'positions':len(positions['positions']),'people':len(rows),'names':name_package,'geography':geo_package,
            'assignments':assignments,'excel_hash':sha256(excel)}


def verify(engine, prepared):
    require_current_schema(engine)
    with engine.connect() as db:
        if db.exec_driver_sql('PRAGMA integrity_check').fetchall()!=[('ok',)]: raise ValueError('integrity_check не пройдено.')
        if db.exec_driver_sql('PRAGMA foreign_key_check').fetchall(): raise ValueError('foreign_key_check не пройдено.')
        from app.repositories.seed_repository import database_snapshot
        result,names_and_ranks,invalid_sex,people=database_snapshot(db)
        if result['person_assignments']!=len(prepared['assignments']): raise ValueError('Неповний імпорт тестових призначень.')
        if result['name_dictionaries']!=prepared['names']['count'] or result['persons']!=prepared['people'] or result['positions']!=prepared['positions']:
            raise ValueError('Кількість записів після засівання не відповідає preflight.')
        for kind,table in [('country','countries'),('region','regions'),('district','districts'),('community','communities'),('settlement','settlements'),('street','streets')]:
            actual=result[table]
            if actual!=prepared['geography']['counts'][table]: raise ValueError(f'{table}: неповний імпорт.')
            result[table]=actual
        if result['geo_city_districts']!=prepared['geography']['counts']['city_districts']: raise ValueError('Неповний імпорт city_districts.')
        if any(number!=result['persons'] for number in names_and_ranks):
            raise ValueError('Не кожна людина має єдиний активний ПІБ/звання.')
        if invalid_sex: raise ValueError('Не задана стать тестової людини.')
        from app.seeds.seed_event_types import EVENT_TYPES
        from app.seeds.seed_event_subtypes import EVENT_SUBTYPES
        if result['ranks']!=len(ReferenceRanks.RANKS) or result['event_types']!=len(EVENT_TYPES) or result['event_subtypes']!=sum(map(len,EVENT_SUBTYPES.values())):
            raise ValueError('Неповний імпорт звань або довідників подій.')
    factory=sessionmaker(bind=engine)
    from app.services.geography import page
    with factory() as db:
        for key in people: page(db,key)
    return result


def populate(path, excel, names, geography, prepared, *, progress=print):
    if sha256(excel)!=prepared['excel_hash']: raise ValueError('Excel змінився після preflight.')
    # Recheck raw package content before writes; CRLF-only Git normalization is explicitly accepted.
    check_hashes(names); check_hashes(geography)
    if sha256(Path(names)/'SHA256SUMS.txt')!=prepared['names']['digest'] or sha256(Path(geography)/'SHA256SUMS.txt')!=prepared['geography']['digest']:
        raise ValueError('Пакет довідників змінився після preflight.')
    engine=create_engine(URL.create('sqlite',database=str(Path(path).resolve())))
    factory=sessionmaker(bind=engine,autoflush=False)
    timings={}
    try:
        init_db(engine)
        with factory() as db: seed_all_structure(db)  # References only, no implicit staff dates.
        started=monotonic(); progress('Імпорт словників ПІБ…')
        name_report=seed_name_package(factory,names,prepared['names']); timings['names']=round(monotonic()-started,2)
        started=monotonic(); progress('Імпорт географії пакетами…')
        geo_report=seed_geography(engine,geography,prepared['geography']); timings['geography']=round(monotonic()-started,2)
        with factory() as db:
            from app.seeds.seed_reference import seed_positions
            with db.begin(): version_id=seed_positions(db,excel,**prepared['config'])
        people=seed_test_people(factory,excel)
        if people.issues: raise ValueError('; '.join(f'persons, рядок {r.row_number}: {r.message}' for r in people.issues))
        from app.services.seed_assignments import seed_assignments
        seed_assignments(factory,prepared['assignments'],version_id)
        if sha256(excel)!=prepared['excel_hash']: raise ValueError('Excel змінився під час імпорту.')
        check_hashes(names); check_hashes(geography)
        result=verify(engine,prepared)
        result.update(timings=timings,created_people=people.created,skipped_people=people.skipped,
                      name_import=name_report,geo_import=geo_report,excluded=prepared['geography']['excluded'],errors=0,duplicates=0)
        return result
    finally: engine.dispose()
